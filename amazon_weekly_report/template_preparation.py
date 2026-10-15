from __future__ import annotations

import re
from copy import copy, deepcopy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import column_index_from_string

from .template_inspection import DAY_PATTERN, TemplateInspectionError, inspect_template


INVALID_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")


class TemplatePreparationError(Exception):
    """无法在不丢失模板内容的前提下准备 SKU 工作表。"""


def _validate_skus(skus: list[str]) -> list[str]:
    normalized = [sku.strip() for sku in skus if sku.strip()]
    if not normalized:
        raise TemplatePreparationError("至少输入一个 SKU")
    seen: set[str] = set()
    for sku in normalized:
        if len(sku) > 31:
            raise TemplatePreparationError(f"SKU 超过 Excel 工作表名称 31 字符限制：{sku}")
        if INVALID_SHEET_CHARACTERS.search(sku) or any(ord(character) < 32 for character in sku):
            raise TemplatePreparationError(f"SKU 含 Excel 工作表名称不支持的字符：{sku}")
        if sku.casefold() in seen:
            raise TemplatePreparationError(f"SKU 重复（忽略大小写）：{sku}")
        seen.add(sku.casefold())
    return normalized


def _copy_template_features(source, target) -> None:
    # openpyxl's copy_worksheet omits conditional formatting and data validation.
    target.views = deepcopy(source.views)
    target.data_validations = deepcopy(source.data_validations)
    target.conditional_formatting = deepcopy(source.conditional_formatting)
    target.protection = copy(source.protection)
    target.auto_filter = deepcopy(source.auto_filter)
    for attribute in ("_print_area", "_print_rows", "_print_cols"):
        setattr(target, attribute, deepcopy(getattr(source, attribute)))
    for attribute in (
        "oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"
    ):
        setattr(target, attribute, copy(getattr(source, attribute)))


def _clear_example_values(worksheet, day_rows: list[int], sku: str) -> None:
    # Clear product-specific metadata while preserving styles, merges, and formulas.
    for row in (2, 3):
        for column in range(1, 9):
            cell = worksheet.cell(row=row, column=column)
            if isinstance(cell, MergedCell):
                continue
            if cell.data_type != "f":
                cell.value = None
    if worksheet["B2"].data_type != "f":
        worksheet["B2"] = sku

    # Retain day/week labels; clear copied sample dates and metrics, not formulas.
    for row in day_rows:
        for column in range(3, worksheet.max_column + 1):
            cell = worksheet.cell(row=row, column=column)
            if cell.data_type != "f":
                cell.value = None


def prepare_sku_workbook(template_path: Path, output_path: Path, skus: list[str]) -> Path:
    """Create clean SKU tabs from a format sample; no operational metrics are filled."""
    template_path = template_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()
    skus = _validate_skus(skus)
    if output_path.suffix.lower() != ".xlsx":
        raise TemplatePreparationError("输出文件必须使用 .xlsx 扩展名")
    if template_path == output_path:
        raise TemplatePreparationError("输出路径不能覆盖原模板")
    if output_path.exists():
        raise TemplatePreparationError(f"输出文件已存在，为保护已有周报不覆盖：{output_path}")

    try:
        summary = inspect_template(template_path)
        workbook = load_workbook(template_path, data_only=False)
    except (OSError, TemplateInspectionError) as exc:
        raise TemplatePreparationError(str(exc)) from exc

    try:
        source = workbook[summary.sheet_name]
        if source._images or source._charts or source.tables:
            raise TemplatePreparationError(
                "模板含图片、图表或 Excel 表格对象；为避免复制时丢失，请先确认如何处理这些对象"
            )

        day_column_number = column_index_from_string(summary.day_column)
        collisions = {name.casefold() for name in workbook.sheetnames if name != source.title}
        conflicting = [sku for sku in skus if sku.casefold() in collisions]
        if conflicting:
            raise TemplatePreparationError(
                "SKU 工作表名称与模板中的其他工作表重名：" + ", ".join(conflicting)
            )

        day_rows = [
            row
            for row in range(summary.header_row + 1, source.max_row + 1)
            if DAY_PATTERN.fullmatch(str(source.cell(row, day_column_number).value or "").strip())
        ]
        if len(day_rows) != summary.day_count:
            raise TemplatePreparationError("模板日期行在读取时发生变化，停止生成")

        for sku in skus:
            worksheet = workbook.copy_worksheet(source)
            _copy_template_features(source, worksheet)
            worksheet.title = sku
            _clear_example_values(worksheet, day_rows, sku)

        workbook.remove(source)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
    except Exception:
        workbook.close()
        raise
    else:
        workbook.close()
    return output_path
