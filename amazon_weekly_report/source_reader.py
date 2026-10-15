from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class SourceReadError(Exception):
    """来源报表无法按明确的字段映射读取。"""


@dataclass(frozen=True)
class MappedRows:
    path: str
    report_kind: str
    rows: list[dict[str, Any]]
    source_encoding: str | None = None


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_csv_rows(path: Path) -> tuple[list[list[Any]], str]:
    raw = path.read_bytes()
    decoded: str | None = None
    encoding_used: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = raw.decode(encoding)
            encoding_used = encoding
            break
        except UnicodeDecodeError:
            continue
    if decoded is None or encoding_used is None:
        raise SourceReadError(f"无法识别 CSV 编码：{path.name}")
    try:
        dialect = csv.Sniffer().sniff(decoded[:64_000], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(decoded), dialect)), encoding_used


def _display_sku(value: Any, number_format: str = "General") -> str | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        fmt = number_format.split(";")[0]
        if fmt and set(fmt) == {"0"}:
            return str(int(value)).zfill(len(fmt))
    return _text(value)


def read_mapped_table(
    path: Path,
    *,
    report_kind: str,
    header_row: int,
    source_headers: dict[str, str],
    sheet_name: str | None = None,
) -> MappedRows:
    """Read exact, user-configured source headers without guessing or imputing values.

    Values stay in their source types. Blank cells become None; explicit numeric
    zero remains zero. SKU text and Excel zero-padding formats are preserved.
    """
    path = path.expanduser().resolve()
    if not path.is_file():
        raise SourceReadError(f"找不到来源文件：{path}")
    if header_row < 1:
        raise SourceReadError("header_row 必须从 1 开始")
    if not source_headers:
        raise SourceReadError("source_headers 为空，拒绝猜测源表字段")

    encoding: str | None = None
    workbook = None
    if path.suffix.lower() == ".csv":
        raw_rows, encoding = _read_csv_rows(path)
        if len(raw_rows) < header_row:
            raise SourceReadError(f"表头行 {header_row} 超出文件行数")
        header_values = raw_rows[header_row - 1]
        body_rows = raw_rows[header_row:]
        row_values = [(values, {}) for values in body_rows]
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
        except Exception as exc:
            raise SourceReadError(f"无法读取工作簿 {path.name}：{exc}") from exc
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                workbook.close()
                raise SourceReadError(f"{path.name} 中找不到工作表：{sheet_name}")
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook[workbook.sheetnames[0]]
        rows_iter = worksheet.iter_rows(min_row=1, values_only=False)
        for _ in range(header_row - 1):
            next(rows_iter, None)
        header_cells = next(rows_iter, None)
        if header_cells is None:
            workbook.close()
            raise SourceReadError(f"表头行 {header_row} 超出工作表行数")
        header_values = [cell.value for cell in header_cells]
        body_rows = []
        row_values = []
        for cells in rows_iter:
            row_values.append(([cell.value for cell in cells], {i: cell.number_format for i, cell in enumerate(cells)}))
    else:
        if path.suffix.lower() == ".xls":
            raise SourceReadError(f"暂不支持旧版 XLS：{path.name}；请另存为 XLSX")
        raise SourceReadError(f"不支持的文件类型：{path.suffix or '(无扩展名)'}")

    if workbook is not None:
        workbook.close()

    headers = [_text(value) for value in header_values]
    positions: dict[str, int] = {}
    for canonical_field, source_header in source_headers.items():
        if not isinstance(source_header, str) or not source_header.strip():
            continue
        matches = [index for index, header in enumerate(headers) if header == source_header.strip()]
        if not matches:
            raise SourceReadError(f"{path.name} 缺少已配置源表头：{source_header!r}（字段 {canonical_field}）")
        if len(matches) > 1:
            raise SourceReadError(f"{path.name} 中源表头重复，无法唯一映射：{source_header!r}")
        positions[canonical_field] = matches[0]
    if not positions:
        raise SourceReadError("没有可用的字段映射")

    records: list[dict[str, Any]] = []
    for values, number_formats in row_values:
        if not any(_text(value) for value in values):
            continue
        record: dict[str, Any] = {}
        for field, position in positions.items():
            value = values[position] if position < len(values) else None
            if field == "sku":
                record[field] = _display_sku(value, number_formats.get(position, "General"))
            else:
                record[field] = None if value is None or value == "" else value
        records.append(record)

    return MappedRows(str(path), report_kind, records, encoding)
