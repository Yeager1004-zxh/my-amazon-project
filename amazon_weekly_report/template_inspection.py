from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DAY_PATTERN = re.compile(r"^D(\d+)$", re.IGNORECASE)


class TemplateInspectionError(Exception):
    """周报模板无法读取或无法识别。"""


@dataclass(frozen=True)
class TemplateInspection:
    path: str
    sheet_names: list[str]
    sheet_name: str
    header_row: int
    first_day_row: int
    last_day_row: int
    day_count: int
    day_column: str
    week_start_rows: list[int]
    columns: dict[str, str]
    manual_input_cells: dict[str, str]
    merged_ranges: list[str]
    formula_count: int
    existing_sku: str | None
    warnings: list[str]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _clean(value: Any) -> str:
    return " ".join(str(value).replace("\n", " ").split()) if value is not None else ""


def inspect_template(path: Path) -> TemplateInspection:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise TemplateInspectionError(f"找不到模板文件：{path}")
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        raise TemplateInspectionError(f"无法读取模板：{exc}") from exc

    try:
        matches: list[tuple[Any, int, dict[str, int]]] = []
        for worksheet in workbook.worksheets:
            for row_number in range(1, min(worksheet.max_row, 30) + 1):
                values = {
                    _clean(cell.value): cell.column
                    for cell in worksheet[row_number]
                    if _clean(cell.value)
                }
                if "天数" in values and "周开始日期" in values:
                    matches.append((worksheet, row_number, values))
        if len(matches) != 1:
            raise TemplateInspectionError(
                f"需要恰好识别到一个包含“天数”和“周开始日期”的表头行，实际找到 {len(matches)} 个"
            )

        worksheet, header_row, headers = matches[0]
        day_column_number = headers["天数"]
        day_column = worksheet.cell(1, day_column_number).column_letter
        day_rows: list[tuple[int, int]] = []
        week_rows: list[int] = []
        for row_number in range(header_row + 1, worksheet.max_row + 1):
            day_value = _clean(worksheet.cell(row_number, day_column_number).value)
            day_match = DAY_PATTERN.fullmatch(day_value)
            if day_match:
                day_rows.append((row_number, int(day_match.group(1))))
            week_value = _clean(worksheet.cell(row_number, 1).value)
            if re.fullmatch(r"第\s*\d+\s*周", week_value):
                week_rows.append(row_number)

        if not day_rows:
            raise TemplateInspectionError("表头下方没有识别到 D01、D02 等日期行")

        formulas = sum(
            1
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        )
        manual_cells: dict[str, str] = {}
        # The protected fields are explicitly named in the user's requirements.
        if _clean(worksheet["A1"].value) == "推广阶段":
            manual_cells["推广阶段"] = "A2"
        if _clean(worksheet["G1"].value) == "发货时间":
            manual_cells["发货时间"] = "G2"
        target_column = headers.get("日目标销量")
        if target_column:
            first_day_row, last_day_row = day_rows[0][0], day_rows[-1][0]
            column_letter = worksheet.cell(1, target_column).column_letter
            manual_cells["日目标销量"] = f"{column_letter}{first_day_row}:{column_letter}{last_day_row}"

        sku = _clean(worksheet["B2"].value) or None
        warnings: list[str] = []
        expected_days = list(range(1, len(day_rows) + 1))
        actual_days = [day for _, day in day_rows]
        if actual_days != expected_days:
            warnings.append("D 编号存在缺号或顺序异常")
        if formulas == 0:
            warnings.append("工作表未发现公式；现有计算结果以静态值保存")
        if sku:
            warnings.append(f"模板样例含 SKU {sku} 和静态业务数据；新建 SKU 页签时应清空样例指标")

        return TemplateInspection(
            path=str(path),
            sheet_names=workbook.sheetnames,
            sheet_name=worksheet.title,
            header_row=header_row,
            first_day_row=day_rows[0][0],
            last_day_row=day_rows[-1][0],
            day_count=len(day_rows),
            day_column=day_column,
            week_start_rows=week_rows,
            columns={name: worksheet.cell(header_row, column).column_letter for name, column in headers.items()},
            manual_input_cells=manual_cells,
            merged_ranges=[str(item) for item in worksheet.merged_cells.ranges],
            formula_count=formulas,
            existing_sku=sku,
            warnings=warnings,
        )
    finally:
        workbook.close()
