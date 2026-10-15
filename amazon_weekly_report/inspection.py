from __future__ import annotations

import csv
import io
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 30


class InspectionError(Exception):
    """输入目录无法检查。"""


@dataclass(frozen=True)
class TableInspection:
    path: str
    sheet_name: str | None
    header_row: int | None
    headers: list[str]
    encoding: str | None = None
    delimiter: str | None = None
    warnings: list[str] | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class FolderInspection:
    directory: str
    files: list[TableInspection]
    unsupported_files: list[str]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _candidate_header(rows: list[tuple[int, list[Any]]]) -> tuple[int | None, list[str]]:
    candidates: list[tuple[tuple[int, int, int], int, list[str]]] = []
    for row_number, values in rows:
        texts = [_cell_text(value) for value in values]
        nonempty = [text for text in texts if text]
        if not nonempty:
            continue
        unique_count = len(set(nonempty))
        # Prefer a fuller row, mostly text cells, then the earliest row. This is
        # only a candidate for review; it never selects or assigns data fields.
        text_count = sum(isinstance(value, str) for value in values if value is not None)
        score = (len(nonempty), text_count, -row_number)
        candidates.append((score, row_number, texts))
    if not candidates:
        return None, []
    _, row_number, values = max(candidates, key=lambda item: item[0])
    return row_number, values


def _header_warnings(headers: list[str]) -> list[str]:
    warnings: list[str] = []
    nonempty = [name for name in headers if name]
    duplicates = sorted(name for name, count in Counter(nonempty).items() if count > 1)
    if duplicates:
        warnings.append("候选表头存在重名列：" + ", ".join(duplicates))
    if any(not name for name in headers):
        warnings.append("候选表头存在空列名")
    return warnings


def _read_csv(path: Path) -> TableInspection:
    raw = path.read_bytes()[:256_000]
    decoded: str | None = None
    encoding_used: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            decoded = raw.decode(encoding)
            encoding_used = encoding
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise InspectionError(f"无法识别 CSV 编码：{path.name}（尝试 UTF-8 和 GB18030）")

    try:
        dialect = csv.Sniffer().sniff(decoded[:32_000], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(decoded), dialect)
    rows: list[tuple[int, list[Any]]] = []
    for row_number, row in enumerate(reader, start=1):
        if any(_cell_text(value) for value in row):
            rows.append((row_number, row))
        if row_number >= HEADER_SCAN_ROWS:
            break
    header_row, headers = _candidate_header(rows)
    return TableInspection(
        path=str(path),
        sheet_name=None,
        header_row=header_row,
        headers=headers,
        encoding=encoding_used,
        delimiter=dialect.delimiter,
        warnings=_header_warnings(headers),
    )


def _read_workbook(path: Path) -> list[TableInspection]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    except Exception as exc:
        raise InspectionError(f"无法读取工作簿 {path.name}：{exc}") from exc

    results: list[TableInspection] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[tuple[int, list[Any]]] = []
            for row_number, cells in enumerate(
                worksheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1
            ):
                values = list(cells)
                if any(_cell_text(value) for value in values):
                    rows.append((row_number, values))
            header_row, headers = _candidate_header(rows)
            if header_row is None:
                continue
            results.append(
                TableInspection(
                    path=str(path),
                    sheet_name=worksheet.title,
                    header_row=header_row,
                    headers=headers,
                    warnings=_header_warnings(headers),
                )
            )
    finally:
        workbook.close()
    return results


def inspect_folder(folder: Path) -> FolderInspection:
    folder = folder.expanduser().resolve()
    if not folder.exists() or not folder.is_dir():
        raise InspectionError(f"目录不存在或不是文件夹：{folder}")

    files: list[TableInspection] = []
    unsupported: list[str] = []
    candidates = sorted(path for path in folder.rglob("*") if path.is_file() and not path.name.startswith("~$"))
    for path in candidates:
        suffix = path.suffix.lower()
        if suffix == ".xls":
            unsupported.append(str(path))
        elif suffix == ".csv":
            files.append(_read_csv(path))
        elif suffix in {".xlsx", ".xlsm"}:
            files.extend(_read_workbook(path))

    return FolderInspection(str(folder), files, unsupported)
