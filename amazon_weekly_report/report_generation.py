from __future__ import annotations

import os
import tempfile
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from .audit import write_validation_log
from .mapping import load_mapping, validate_mapping
from .source_reader import MappedRows, SourceReadError, read_mapped_table
from .template_expansion import TemplateExpansionError, ensure_day_rows, write_timeline_rows
from .template_inspection import TemplateInspection, TemplateInspectionError, inspect_template
from .template_preparation import (
    TemplatePreparationError,
    _clear_example_values,
    _copy_template_features,
    _validate_skus,
    prepare_sku_workbook,
)
from .timeline import LaunchDay, build_timeline
from .providers.lingxing_api import MAX_PRODUCT_QUERY_DAYS, SkuDailySalesWindow


class ReportGenerationError(RuntimeError):
    """Report generation stopped to prevent ambiguous or fabricated data."""


@dataclass(frozen=True)
class SkuGenerationResult:
    sku: str
    status: str
    d01: str | None
    d01_source: str | None
    latest_covered_date: str | None
    days: int
    sales_units_total: float | None
    sales_amount_total: float | str | None
    actual_ad_spend_total: float | None
    warnings: list[str]


@dataclass(frozen=True)
class GenerationResult:
    output_path: Path
    audit_path: Path
    sku_results: list[SkuGenerationResult]


@dataclass
class ReportData:
    name: str
    rows_by_sku_date: dict[str, dict[date, dict[str, Any]]]
    coverage: dict[str, set[date]]
    source_files: list[str]
    source_total_spend: Decimal | None = None
    accepted_total_spend: Decimal | None = None
    accepted_spend_by_sku_date: dict[str, dict[date, Decimal]] | None = None
    all_skus: set[str] | None = None
    warnings: list[str] | None = None
    anomalies: list[str] | None = None
    orders_by_sku_date: dict[str, dict[date, Decimal | None]] | None = None


_NUMERIC_FIELDS = {
    "units_sold", "sales_amount", "unit_price", "organic_orders", "actual_spend",
    "ad_orders", "ad_sales", "impressions", "clicks", "sessions", "organic_clicks",
}


def _parse_date(value: Any, date_format: str | None) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str) or not value.strip():
        return None
    text = value.strip()
    if date_format:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        # Do not guess between US and international slash-date formats.
        return None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = text.replace(",", "").replace("$", "").strip()
    try:
        parsed = Decimal(text)
    except InvalidOperation:
        return None
    if not parsed.is_finite():
        return None
    return -parsed if negative else parsed


def _excel_number(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _safe_pattern(folder: Path, pattern: str) -> list[Path]:
    relative = Path(pattern)
    if relative.is_absolute() or ".." in relative.parts:
        raise ReportGenerationError("file_patterns 只能使用数据目录内的相对文件名模式。")
    matches = sorted(
        path.resolve()
        for path in folder.glob(pattern)
        if path.is_file() and path.suffix.lower() in {".csv", ".xlsx", ".xlsm"} and not path.name.startswith("~$")
    )
    return matches


def _report_is_configured(name: str, config: dict[str, Any]) -> bool:
    headers = config.get("source_headers", {})
    return bool(config.get("file_patterns")) or any(
        isinstance(value, str) and value.strip() for value in headers.values()
    )


def _load_report(
    folder: Path,
    report_name: str,
    config: dict[str, Any],
    *,
    required: bool,
) -> ReportData:
    rows_by_sku_date: dict[str, dict[date, dict[str, Any]]] = defaultdict(dict)
    coverage: dict[str, set[date]] = defaultdict(set)
    warnings: list[str] = []
    anomalies: list[str] = []
    files: set[Path] = set()
    for pattern in config.get("file_patterns", []):
        files.update(_safe_pattern(folder, pattern))
    if not files:
        if required:
            raise ReportGenerationError(f"未找到 {report_name} 配置指定的输入文件。")
        return ReportData(report_name, {}, {}, [], warnings=[f"没有找到 {report_name} 来源文件；相应单元格保持空白。"], anomalies=[])

    source_headers = config.get("source_headers", {})
    all_rows: list[dict[str, Any]] = []
    for file_path in sorted(files):
        try:
            mapped: MappedRows = read_mapped_table(
                file_path,
                report_kind=report_name,
                header_row=int(config["header_row"]),
                source_headers=source_headers,
                sheet_name=config.get("sheet_name") or None,
            )
        except (OSError, ValueError, SourceReadError, KeyError, TypeError) as exc:
            raise ReportGenerationError(f"{report_name} 来源 {file_path.name} 无法按映射读取：{exc}") from exc
        for row_number, raw in enumerate(mapped.rows, start=1):
            row = dict(raw)
            row["_source_file"] = file_path.name
            row["_source_row"] = row_number + int(config["header_row"])
            sku = str(row.get("sku") or "").strip()
            store = str(row.get("store") or "").strip()
            day = _parse_date(row.get("date"), config.get("date_format"))
            if not sku or not store:
                anomalies.append(f"{file_path.name} 第 {row['_source_row']} 行缺少 SKU 或店铺，已排除。")
                continue
            if report_name != "product_snapshot" and day is None:
                anomalies.append(
                    f"{file_path.name} 第 {row['_source_row']} 行日期为空或格式不明确，已排除；请在映射配置中设置 date_format。"
                )
                continue
            row["_sku_key"] = sku.casefold()
            row["_store"] = store
            row["_date"] = day
            invalid_numbers = []
            for field in _NUMERIC_FIELDS:
                if field not in row or row[field] is None:
                    continue
                parsed = _parse_decimal(row[field])
                if parsed is None:
                    invalid_numbers.append(field)
                else:
                    row[field] = parsed
            if invalid_numbers:
                anomalies.append(
                    f"{file_path.name} 第 {row['_source_row']} 行的数值字段无法解析（{', '.join(invalid_numbers)}），该行已排除。"
                )
                continue
            if "actual_spend" in row and isinstance(row["actual_spend"], Decimal) and row["actual_spend"] < 0:
                anomalies.append(f"{file_path.name} 第 {row['_source_row']} 行广告实际花费为负数，已排除。")
                continue
            all_rows.append(row)
            if day is not None:
                coverage[row["_sku_key"]].add(day)

    all_skus = {row["_sku_key"] for row in all_rows}
    spend_source_rows = [row["actual_spend"] for row in all_rows if isinstance(row.get("actual_spend"), Decimal)]
    report_total = sum(spend_source_rows, Decimal("0")) if spend_source_rows else None

    campaign_field = "campaign_id" if any(row.get("campaign_id") not in (None, "") for row in all_rows) else None
    ad_group_field = "ad_group_id" if any(row.get("ad_group_id") not in (None, "") for row in all_rows) else None
    duplicates: set[tuple[Any, ...]] = set()
    first_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in all_rows:
        duplicate_key = (row["_store"], row["_sku_key"], row["_date"])
        if report_name == "actual_ad_spend" and campaign_field:
            duplicate_key += (str(row.get("campaign_id") or "").strip(),)
        if report_name == "actual_ad_spend" and ad_group_field:
            duplicate_key += (str(row.get("ad_group_id") or "").strip(),)
        if report_name == "ad_performance" and campaign_field:
            duplicate_key += (str(row.get("campaign_id") or "").strip(),)
        if report_name == "ad_performance" and ad_group_field:
            duplicate_key += (str(row.get("ad_group_id") or "").strip(),)
        if duplicate_key in first_by_key:
            duplicates.add(duplicate_key)
        else:
            first_by_key[duplicate_key] = row
    if duplicates:
        anomalies.append(
            f"{report_name} 有 {len(duplicates)} 组重复的店铺/SKU/日期记录，无法区分重复导出与真实明细；相关记录均未累计。"
        )

    clean_rows = [row for key, row in first_by_key.items() if key not in duplicates]
    if report_name == "actual_ad_spend" and (campaign_field or ad_group_field):
        # Campaign/group-level cost is often repeated once per advertised SKU.
        # An identical amount with the same campaign scope and day cannot be
        # safely attributed to each SKU, so exclude every row in that signature.
        spend_by_campaign_signature: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
        for row in clean_rows:
            spend = row.get("actual_spend")
            if not isinstance(spend, Decimal):
                continue
            signature = (
                row["_store"], row["_date"],
                str(row.get("campaign_id") or "").strip(),
                str(row.get("ad_group_id") or "").strip(),
                spend,
            )
            spend_by_campaign_signature[signature].append(row)
        cross_sku_duplicates = {
            signature
            for signature, matching_rows in spend_by_campaign_signature.items()
            if len({row["_sku_key"] for row in matching_rows}) > 1
        }
        if cross_sku_duplicates:
            affected_skus = sorted({
                row["sku"]
                for signature in cross_sku_duplicates
                for row in spend_by_campaign_signature[signature]
            })
            anomalies.append(
                f"actual_ad_spend 有 {len(cross_sku_duplicates)} 组相同活动/广告组花费重复关联多个 SKU "
                f"（{', '.join(affected_skus)}）；相关记录均未归属或累计，请核对领星原始报表。"
            )
            clean_rows = [
                row for row in clean_rows
                if (
                    row["_store"], row["_date"],
                    str(row.get("campaign_id") or "").strip(),
                    str(row.get("ad_group_id") or "").strip(),
                    row.get("actual_spend"),
                ) not in cross_sku_duplicates
            ]
    accepted_total = None
    if report_name == "actual_ad_spend":
        accepted_spend_rows = [row["actual_spend"] for row in clean_rows if isinstance(row.get("actual_spend"), Decimal)]
        accepted_total = sum(accepted_spend_rows, Decimal("0")) if accepted_spend_rows else None
    accepted_spend_by_sku_date: dict[str, dict[date, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    if report_name == "actual_ad_spend":
        for row in clean_rows:
            if isinstance(row.get("actual_spend"), Decimal) and row.get("_date") is not None:
                accepted_spend_by_sku_date[row["_sku_key"]][row["_date"]] += row["actual_spend"]

    requested_stores: dict[str, set[str]] = defaultdict(set)
    for row in all_rows:
        requested_stores[row["_sku_key"]].add(row["_store"])
    by_sku_date_store: dict[tuple[str, date, str], list[dict[str, Any]]] = defaultdict(list)
    for row in clean_rows:
        if row["_date"] is None:
            continue
        by_sku_date_store[(row["_sku_key"], row["_date"], row["_store"])].append(row)

    by_sku_date: dict[tuple[str, date], list[dict[str, Any]]] = defaultdict(list)
    for (sku_key, day, _store), store_rows in by_sku_date_store.items():
        by_sku_date[(sku_key, day)].extend(store_rows)
    daily: dict[str, dict[date, dict[str, Any]]] = defaultdict(dict)
    for (sku_key, day), rows in by_sku_date.items():
        stores_for_sku = requested_stores[sku_key]
        present_stores = {row["_store"] for row in rows}
        if len(stores_for_sku) > 1 and present_stores != stores_for_sku:
            warnings.append(
                f"{report_name} {day.isoformat()} 的 {sku_key} 未覆盖所有出现过的店铺，合计字段保持空白。"
            )
            continue
        fields = set().union(*(row.keys() for row in rows))
        values: dict[str, Any] = {}
        for field in _NUMERIC_FIELDS:
            present = [row.get(field) for row in rows]
            if any(value is None for value in present):
                values[field] = None
            elif present and all(isinstance(value, Decimal) for value in present):
                values[field] = sum(present, Decimal("0"))
        if report_name == "actual_ad_spend":
            for row in rows:
                spend = row.get("actual_spend")
                disabled_at = _parse_date(row.get("campaign_disabled_at"), config.get("date_format"))
                if isinstance(spend, Decimal) and spend > 0 and disabled_at and day > disabled_at:
                    campaign = str(row.get("campaign_name") or row.get("campaign_id") or "(未映射活动名)")
                    warnings.append(
                        f"检测到活动 {campaign} 的关闭日期 {disabled_at.isoformat()} 后仍有实际花费记录 {day.isoformat()}；请人工核对原始报表。"
                    )
        if fields:
            daily[sku_key][day] = values

    return ReportData(
        name=report_name,
        rows_by_sku_date=dict(daily),
        coverage=dict(coverage),
        source_files=[path.name for path in sorted(files)],
        source_total_spend=report_total,
        accepted_total_spend=accepted_total,
        accepted_spend_by_sku_date={sku: dict(days) for sku, days in accepted_spend_by_sku_date.items()},
        all_skus=all_skus,
        warnings=warnings,
        anomalies=anomalies,
    )


def _aggregate_metric(reports: dict[str, ReportData], sku_key: str, day: date, field: str) -> Decimal | None:
    for report in reports.values():
        value = report.rows_by_sku_date.get(sku_key, {}).get(day, {}).get(field)
        if value is not None:
            return value
    return None


def _ratio(numerator: Decimal | None, denominator: Decimal | None) -> Decimal | None:
    if numerator is None or denominator is None or denominator == 0:
        return None
    return numerator / denominator


def _manual_number(value: Any) -> Decimal | None:
    return _parse_decimal(value)


def _header_columns(summary: TemplateInspection) -> dict[str, int]:
    from openpyxl.utils.cell import column_index_from_string

    aliases = {
        "week": "周次", "day": "天数", "week_date": "周开始日期", "price": "售价",
        "sales": "日销", "organic_orders": "自然单", "ad_orders": "广告单",
        "ad_order_share": "广告订单 占比", "category": "类目", "rank": "小类排名",
        "rating": "评分", "reviews": "评论数", "daily_target": "日目标销量",
        "target_attainment": "目标达成率", "sessions": "Sessions", "organic_clicks": "自然点击",
        "impressions": "广告曝光", "ad_clicks": "广告点击", "ctr": "广告CTR", "ad_cvr": "广告CVR",
        "spend": "广告花费 ($)", "ad_sales": "广告销售额 ($)", "overall_cvr": "整体CVR", "acos": "ACOS",
    }
    output: dict[str, int] = {}
    missing = []
    for key, header in aliases.items():
        letter = summary.columns.get(header)
        if letter:
            output[key] = column_index_from_string(letter)
        elif key in {"week", "day", "week_date", "sales", "daily_target"}:
            missing.append(header)
    if missing:
        raise ReportGenerationError("模板缺少必需字段：" + ", ".join(missing))
    return output


def _write_metric(worksheet, row: int, column: int | None, value: Decimal | int | float | str | None, *, preserve: bool) -> None:
    if column is None or value is None:
        return
    cell = worksheet.cell(row, column)
    if cell.data_type == "f":
        return
    if preserve and cell.value not in (None, ""):
        return
    if isinstance(value, Decimal):
        cell.value = _excel_number(value)
    else:
        cell.value = value


def _clear_or_copy_sku_sheet(workbook, source, summary: TemplateInspection, sku: str):
    if source._images or source._charts or source.tables:
        raise ReportGenerationError("模板含图片、图表或 Excel 表格对象；为避免丢失，已停止生成。")
    worksheet = workbook.copy_worksheet(source)
    _copy_template_features(source, worksheet)
    worksheet.title = sku
    day_columns = column_index_from_string(summary.day_column)
    day_rows = [
        row for row in range(summary.header_row + 1, source.max_row + 1)
        if isinstance(source.cell(row, day_columns).value, str)
        and source.cell(row, day_columns).value.upper().startswith("D")
    ]
    _clear_example_values(worksheet, day_rows, sku)
    return worksheet


def _metric_values(
    sku_key: str,
    day: date,
    reports: dict[str, ReportData],
    *,
    target: Any = None,
) -> dict[str, Decimal | None]:
    sales = reports["sales"].rows_by_sku_date.get(sku_key, {}).get(day, {})
    ad = reports.get("ad_performance")
    ad_values = ad.rows_by_sku_date.get(sku_key, {}).get(day, {}) if ad else {}
    traffic = reports.get("traffic")
    traffic_values = traffic.rows_by_sku_date.get(sku_key, {}).get(day, {}) if traffic else {}
    spend_report = reports.get("actual_ad_spend")
    spend_values = spend_report.rows_by_sku_date.get(sku_key, {}).get(day, {}) if spend_report else {}

    units = sales.get("units_sold")
    revenue = sales.get("sales_amount")
    unit_price = sales.get("unit_price")
    if unit_price is None:
        unit_price = _ratio(revenue, units)
    ad_orders = ad_values.get("ad_orders")
    clicks = ad_values.get("clicks")
    impressions = ad_values.get("impressions")
    sessions = traffic_values.get("sessions")
    spend = spend_values.get("actual_spend")
    ad_sales = ad_values.get("ad_sales")
    target_number = _manual_number(target)
    return {
        "price": unit_price,
        "sales": units,
        "organic_orders": sales.get("organic_orders"),
        "ad_orders": ad_orders,
        "ad_order_share": _ratio(ad_orders, units),
        "target_attainment": _ratio(units, target_number),
        "sessions": traffic_values.get("sessions"),
        "organic_clicks": traffic_values.get("organic_clicks"),
        "impressions": impressions,
        "ad_clicks": clicks,
        "ctr": _ratio(clicks, impressions),
        "ad_cvr": _ratio(ad_orders, clicks),
        "spend": spend,
        "ad_sales": ad_sales,
        "overall_cvr": _ratio(units, sessions),
        "acos": _ratio(spend, ad_sales),
        "_sales_amount": revenue,
    }


def _days_in_sheet(worksheet, summary: TemplateInspection) -> tuple[list[int], date | None]:
    day_col = column_index_from_string(summary.day_column)
    rows: list[int] = []
    d01: date | None = None
    for row in range(summary.header_row + 1, worksheet.max_row + 1):
        value = worksheet.cell(row, day_col).value
        if isinstance(value, str) and value.strip().upper().startswith("D"):
            rows.append(row)
            if value.strip().upper() == "D01":
                d01 = _date_from_cell(worksheet.cell(row, column_index_from_string(summary.columns["周开始日期"])).value)
    return rows, d01


def read_existing_d01_dates(
    existing_report_path: Path,
    template_path: Path,
    skus: list[str],
) -> dict[str, date]:
    """Read existing workbook D01 dates for a pre-update confirmation prompt."""
    existing_path = existing_report_path.expanduser().resolve()
    if not existing_path.is_file():
        raise ReportGenerationError("找不到要增量更新的现有周报。")
    try:
        summary = inspect_template(template_path.expanduser().resolve())
        workbook = load_workbook(existing_path, data_only=False, read_only=False)
    except (OSError, TemplateInspectionError, ValueError) as exc:
        raise ReportGenerationError("无法读取现有周报的 D01 日期。") from exc
    dates: dict[str, date] = {}
    try:
        for sku in skus:
            worksheet = next((sheet for sheet in workbook.worksheets if sheet.title.casefold() == sku.casefold()), None)
            if worksheet is None:
                continue
            _rows, d01 = _days_in_sheet(worksheet, summary)
            if d01 is not None:
                dates[sku.casefold()] = d01
    finally:
        workbook.close()
    return dates


def _date_from_cell(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _api_sales_report(
    windows: list[SkuDailySalesWindow],
    *,
    skus: list[str],
    store_sid: int,
) -> ReportData:
    """Convert complete API windows to report rows without treating absence as zero."""
    if not windows:
        raise ReportGenerationError("OpenAPI 没有返回任何完整查询日期段。")
    requested = {sku.casefold(): sku for sku in skus}
    daily_values: dict[str, dict[date, dict[str, Any]]] = defaultdict(dict)
    coverage: dict[str, set[date]] = defaultdict(set)
    orders: dict[str, dict[date, Decimal | None]] = defaultdict(dict)
    currencies_by_sku: dict[str, set[str]] = defaultdict(set)
    expected_next_date: date | None = None

    for window in windows:
        if not isinstance(window, SkuDailySalesWindow):
            raise ReportGenerationError("OpenAPI 历史数据包含未知日期段结构。")
        if (
            window.end_date < window.start_date
            or (window.end_date - window.start_date).days + 1 > MAX_PRODUCT_QUERY_DAYS
        ):
            raise ReportGenerationError("OpenAPI 历史数据日期段超过 92 天或日期无效。")
        if expected_next_date is not None and window.start_date != expected_next_date:
            raise ReportGenerationError("OpenAPI 历史日期段存在重叠或缺口；未生成周报。")
        expected_next_date = window.end_date + timedelta(days=1)

        mapping_by_sku: dict[str, Any] = {}
        for mapping in window.mappings:
            key = mapping.local_sku.casefold()
            if key not in requested:
                raise ReportGenerationError("OpenAPI 商品映射返回了未请求的 SKU；未生成周报。")
            if key in mapping_by_sku:
                raise ReportGenerationError("OpenAPI 同一日期段重复返回 SKU 映射；未生成周报。")
            if not mapping.seller_skus or len(set(mapping.seller_skus)) != len(mapping.seller_skus):
                raise ReportGenerationError(f"SKU {requested[key]} 的 MSKU 映射为空或重复；未生成周报。")
            mapping_by_sku[key] = mapping
        missing_mappings = set(requested) - set(mapping_by_sku)
        if missing_mappings:
            names = [requested[key] for key in sorted(missing_mappings)]
            raise ReportGenerationError("OpenAPI 商品映射不完整：" + ", ".join(names) + "；未生成周报。")

        expected: set[tuple[str, date, str]] = set()
        day = window.start_date
        while day <= window.end_date:
            for key, mapping in mapping_by_sku.items():
                for msku in mapping.seller_skus:
                    expected.add((key, day, msku.casefold()))
                    coverage[key].add(day)
            day += timedelta(days=1)

        indexed: dict[tuple[str, date, str], Any] = {}
        for row in window.rows:
            if row.sid != store_sid:
                raise ReportGenerationError("OpenAPI 销售记录 SID 与所选店铺不一致；未生成周报。")
            key = row.local_sku.casefold()
            row_key = (key, row.event_date, row.seller_sku.casefold())
            if row_key not in expected:
                raise ReportGenerationError("OpenAPI 销售记录无法唯一匹配请求的店铺/SKU/MSKU/日期；未生成周报。")
            if row_key in indexed:
                raise ReportGenerationError("OpenAPI 销售记录存在重复店铺/SKU/MSKU/日期；未生成周报。")
            indexed[row_key] = row
            if row.currency_code:
                currencies_by_sku[key].add(row.currency_code)
        if set(indexed) != expected:
            raise ReportGenerationError("OpenAPI 历史日期段未完整返回所有已映射 MSKU 的日期位置；未生成周报。")

        for key, mapping in mapping_by_sku.items():
            day = window.start_date
            while day <= window.end_date:
                day_rows = [indexed[(key, day, msku.casefold())] for msku in mapping.seller_skus]

                def sum_if_complete(field: str) -> Decimal | None:
                    values = [getattr(row, field) for row in day_rows]
                    if any(value is None for value in values):
                        return None
                    if any(isinstance(value, bool) or not isinstance(value, (int, Decimal)) for value in values):
                        raise ReportGenerationError(f"OpenAPI 销售字段 {field} 类型无效；未生成周报。")
                    return sum((Decimal(value) for value in values), Decimal("0"))

                daily_values[key][day] = {
                    "units_sold": sum_if_complete("units_sold"),
                    "sales_amount": sum_if_complete("sales_amount"),
                }
                orders[key][day] = sum_if_complete("order_count")
                day += timedelta(days=1)

    mixed_currencies = [requested[key] for key, codes in currencies_by_sku.items() if len(codes) > 1]
    if mixed_currencies:
        raise ReportGenerationError("同一 SKU 的销售记录出现多个币种，无法安全合计：" + ", ".join(mixed_currencies))
    return ReportData(
        name="sales",
        rows_by_sku_date=dict(daily_values),
        coverage=dict(coverage),
        source_files=[],
        all_skus=set(requested),
        warnings=["来源为领星 OpenAPI；日期覆盖表示请求成功完成到该站点日期；缺失指标仍保持空白。"],
        anomalies=[],
        orders_by_sku_date=dict(orders),
    )


def generate_reports(
    *,
    template_path: Path,
    data_folder: Path | None,
    mapping_path: Path | None,
    skus: list[str],
    output_path: Path,
    d01_dates: dict[str, date] | None = None,
    d01_confirmed: bool = False,
    confirmed_d01_changes: bool = False,
    existing_report_path: Path | None = None,
    api_sales_windows: list[SkuDailySalesWindow] | None = None,
    api_store_sid: int | None = None,
    api_store_name: str | None = None,
) -> GenerationResult:
    """Build from explicit file mappings or completed, validated OpenAPI windows."""
    template_path = template_path.expanduser().resolve()
    data_folder = data_folder.expanduser().resolve() if data_folder is not None else None
    mapping_path = mapping_path.expanduser().resolve() if mapping_path is not None else None
    output_path = output_path.expanduser().resolve()
    normalized_d01_dates: dict[str, date] = {}
    for raw_sku, day in (d01_dates or {}).items():
        if not isinstance(raw_sku, str) or not raw_sku.strip() or type(day) is not date:
            raise ReportGenerationError("D01 必须是逐 SKU 明确输入的有效 YYYY-MM-DD 日期。")
        key = raw_sku.strip().casefold()
        if key in normalized_d01_dates:
            raise ReportGenerationError("D01 设置中存在重复 SKU。")
        normalized_d01_dates[key] = day
    d01_dates = normalized_d01_dates
    try:
        skus = _validate_skus(skus)
    except TemplatePreparationError as exc:
        raise ReportGenerationError(str(exc)) from exc
    if output_path.suffix.lower() != ".xlsx":
        raise ReportGenerationError("输出文件必须使用 .xlsx 扩展名。")
    if output_path.exists():
        raise ReportGenerationError("输出文件已存在，为保护已有周报不覆盖。")
    audit_target = output_path.with_name(output_path.stem + "_数据校验.jsonl")
    if audit_target.exists():
        raise ReportGenerationError("对应的数据校验日志已存在，为保护已有记录不覆盖。")
    if template_path == output_path or (existing_report_path and existing_report_path.resolve() == output_path):
        raise ReportGenerationError("输出路径不能覆盖模板或现有周报。")
    if not d01_confirmed:
        raise ReportGenerationError("D01 尚未由用户明确确认；未生成周报。")
    requested_skus = {sku.casefold() for sku in skus}
    missing_d01 = [sku for sku in skus if sku.casefold() not in d01_dates]
    extra_d01 = [sku for sku in d01_dates if sku not in requested_skus]
    if missing_d01:
        raise ReportGenerationError("以下 SKU 尚未设置并确认 D01，不能生成周报：" + ", ".join(missing_d01))
    if extra_d01:
        raise ReportGenerationError("D01 设置包含本次未请求的 SKU，停止生成以避免错配。")
    api_mode = api_sales_windows is not None
    if api_mode:
        if isinstance(api_store_sid, bool) or not isinstance(api_store_sid, int):
            raise ReportGenerationError("OpenAPI 周报必须绑定动态店铺 SID。")
        if not isinstance(api_store_name, str) or not api_store_name.strip():
            raise ReportGenerationError("OpenAPI 周报必须保留所选店铺名称。")
    else:
        if data_folder is None or not data_folder.is_dir():
            raise ReportGenerationError("数据文件夹不存在。")
        if mapping_path is None:
            raise ReportGenerationError("离线 Excel/CSV 模式必须选择字段映射 JSON。")
    mapping: dict[str, Any] | None = None
    if not api_mode:
        try:
            mapping = load_mapping(mapping_path)
            issues = validate_mapping(mapping)
        except (OSError, ValueError) as exc:
            raise ReportGenerationError(f"映射配置无法读取：{exc}") from exc
        if issues:
            raise ReportGenerationError("字段映射尚未就绪：\n- " + "\n- ".join(issues))
    try:
        summary = inspect_template(template_path)
    except TemplateInspectionError as exc:
        raise ReportGenerationError(str(exc)) from exc

    reports: dict[str, ReportData] = {}
    if api_mode:
        reports["sales"] = _api_sales_report(api_sales_windows or [], skus=skus, store_sid=api_store_sid)
        reports["actual_ad_spend"] = ReportData(
            "actual_ad_spend", {}, {}, [],
            warnings=["本次 OpenAPI 流程尚未接入已核验的广告活动实际花费；相关单元格保持空白。"],
            anomalies=[],
        )
    else:
        reports_config = mapping["reports"]
        reports["sales"] = _load_report(data_folder, "sales", reports_config["sales"], required=True)
        spend_config = reports_config.get("actual_ad_spend", {})
        if _report_is_configured("actual_ad_spend", spend_config):
            reports["actual_ad_spend"] = _load_report(
                data_folder, "actual_ad_spend", spend_config, required=False
            )
        else:
            reports["actual_ad_spend"] = ReportData(
                "actual_ad_spend", {}, {}, [], warnings=["未配置广告实际花费来源；广告花费单元格保持空白。"], anomalies=[]
            )
        for optional_name in ("ad_performance", "traffic"):
            config = reports_config.get(optional_name, {})
            if config.get("enabled") and _report_is_configured(optional_name, config):
                reports[optional_name] = _load_report(data_folder, optional_name, config, required=False)

    sku_info: dict[str, dict[str, Any]] = {}
    failures: list[SkuGenerationResult] = []
    for sku in skus:
        key = sku.casefold()
        warnings = []
        d01 = d01_dates.get(key)
        d01_source = "用户手动确认的店铺站点 D01 日期"
        coverage_dates = set()
        for report in reports.values():
            coverage_dates.update(report.coverage.get(key, set()))
        latest = max(coverage_dates) if coverage_dates else None
        if latest is None or latest < d01:
            failures.append(SkuGenerationResult(sku, "failed", d01.isoformat(), d01_source, latest.isoformat() if latest else None, 0, None, None, None, ["没有覆盖 D01 及之后日期的有效来源记录，未生成 SKU 页签。"]))
            continue
        timeline = build_timeline(d01, latest)
        sku_info[sku] = {"key": key, "d01": d01, "d01_source": d01_source, "latest": latest, "timeline": timeline, "warnings": warnings}

    if not sku_info:
        details = "; ".join(f"{item.sku}: {'; '.join(item.warnings)}" for item in failures)
        raise ReportGenerationError("没有可生成的 SKU。" + details)

    existing_d01_by_sku: dict[str, date] = {}
    if existing_report_path is not None:
        existing_d01_by_sku = read_existing_d01_dates(existing_report_path, template_path, skus)
    changed_d01 = {
        sku: (existing_d01_by_sku[sku.casefold()], info["d01"])
        for sku, info in sku_info.items()
        if sku.casefold() in existing_d01_by_sku and existing_d01_by_sku[sku.casefold()] != info["d01"]
    }
    if changed_d01 and not confirmed_d01_changes:
        names = ", ".join(changed_d01)
        raise ReportGenerationError(
            f"现有周报的 D01 与本次确认日期不同（{names}）；请先确认日期错位风险后重新生成。"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    stage_fd, stage_name = tempfile.mkstemp(prefix="weekly-report-", suffix=".xlsx", dir=output_path.parent)
    os.close(stage_fd)
    os.unlink(stage_name)
    stage_path = Path(stage_name)
    results: list[SkuGenerationResult] = []
    daily_audit_entries: list[dict[str, Any]] = []
    try:
        if existing_report_path is None:
            try:
                prepare_sku_workbook(template_path, stage_path, list(sku_info))
            except (OSError, TemplatePreparationError) as exc:
                raise ReportGenerationError(f"无法从格式样例创建 SKU 工作表：{exc}") from exc
            workbook = load_workbook(stage_path, data_only=False)
            original_sku_names: set[str] = set()
            source_sheets_to_remove: list[str] = []
        else:
            existing_path = existing_report_path.expanduser().resolve()
            if not existing_path.is_file():
                raise ReportGenerationError("找不到要增量更新的现有周报。")
            workbook = load_workbook(existing_path, data_only=False)
            stage_path.unlink(missing_ok=True)
            workbook.save(stage_path)
            workbook.close()
            workbook = load_workbook(stage_path, data_only=False)
            source_sheets_to_remove = []
            original_sku_names = {name.casefold() for name in workbook.sheetnames}
            source_sheet = next((workbook[name] for name in workbook.sheetnames if name == summary.sheet_name), None)
            if source_sheet is None:
                source_sheet = next(
                    (sheet for sheet in workbook.worksheets if sheet.max_row >= summary.last_day_row), None
                )
            if source_sheet is None:
                raise ReportGenerationError("现有周报中没有可用的样式工作表。")
            collisions = {name.casefold() for name in workbook.sheetnames}
            for sku in sku_info:
                if sku.casefold() not in collisions:
                    _clear_or_copy_sku_sheet(workbook, source_sheet, summary, sku)

        columns = _header_columns(summary)
        successful: list[SkuGenerationResult] = []
        for sku, info in sku_info.items():
            key, d01, latest, timeline = info["key"], info["d01"], info["latest"], info["timeline"]
            worksheet = next((sheet for sheet in workbook.worksheets if sheet.title.casefold() == sku.casefold()), None)
            if worksheet is None:
                raise ReportGenerationError(f"未能创建 SKU 工作表：{sku}")
            existing_rows, existing_d01 = _days_in_sheet(worksheet, summary)
            is_existing = sku.casefold() in original_sku_names
            d01_changed = is_existing and existing_d01 is not None and existing_d01 != d01
            row_count = max(len(existing_rows), len(timeline)) if is_existing else len(timeline)
            output_timeline = timeline
            if d01_changed:
                observed_labels = [
                    str(worksheet.cell(row, column_index_from_string(summary.day_column)).value).strip().upper()
                    for row in existing_rows
                ]
                expected_labels = [f"D{index:02d}" for index in range(1, len(existing_rows) + 1)]
                if observed_labels != expected_labels:
                    raise ReportGenerationError(f"现有周报 {sku} 的日期编号不连续，不能安全修改 D01。")
                output_timeline = build_timeline(d01, d01 + timedelta(days=max(row_count, 1) - 1))
            try:
                row_numbers = ensure_day_rows(worksheet, summary, row_count)
            except TemplateExpansionError as exc:
                raise ReportGenerationError(f"SKU {sku} 扩展日期行失败：{exc}") from exc
            if d01_changed:
                write_timeline_rows(worksheet, summary, output_timeline, row_numbers, clear_after=False)
                generated_fields = (
                    "price", "sales", "organic_orders", "ad_orders", "ad_order_share", "target_attainment",
                    "sessions", "organic_clicks", "impressions", "ad_clicks", "ctr", "ad_cvr", "spend",
                    "ad_sales", "overall_cvr", "acos",
                )
                for row_number in row_numbers:
                    for field in generated_fields:
                        column = columns.get(field)
                        if column is not None:
                            cell = worksheet.cell(row_number, column)
                            if cell.data_type != "f":
                                cell.value = None
            elif is_existing and len(existing_rows) >= len(timeline):
                # Keep the existing date axis and previously saved manual input.
                row_numbers = existing_rows
                for index, row in enumerate(row_numbers[:len(timeline)]):
                    if worksheet.cell(row, column_index_from_string(summary.day_column)).value != timeline[index].day_label:
                        failures.append(SkuGenerationResult(sku, "failed", d01.isoformat(), info["d01_source"], latest.isoformat(), 0, None, None, None,
                                                            ["现有周报日期标签不连续或与 D 编号不匹配；未更新该 SKU。"]))
                        break
                if failures and failures[-1].sku == sku and failures[-1].status == "failed":
                    continue
            else:
                write_timeline_rows(worksheet, summary, timeline, row_numbers, clear_after=not is_existing)

            by_day_row = {
                int(str(worksheet.cell(row, column_index_from_string(summary.day_column)).value)[1:]): row
                for row in row_numbers
                if isinstance(worksheet.cell(row, column_index_from_string(summary.day_column)).value, str)
                and worksheet.cell(row, column_index_from_string(summary.day_column)).value.upper().startswith("D")
            }
            sales_units_total = Decimal("0")
            sales_amount_total = Decimal("0")
            has_sales_units = False
            has_sales_amount = False
            ad_spend_total = Decimal("0")
            has_spend = False
            for item in output_timeline:
                row_number = by_day_row.get(item.day_number)
                if row_number is None:
                    continue
                target_value = worksheet.cell(row_number, columns["daily_target"]).value
                values = _metric_values(key, item.date, reports, target=target_value)
                sales_source_present = item.date in reports["sales"].rows_by_sku_date.get(key, {})
                spend_report = reports.get("actual_ad_spend")
                spend_source_present = bool(spend_report and item.date in spend_report.rows_by_sku_date.get(key, {}))
                metrics = {
                    "price": values["price"], "sales": values["sales"], "organic_orders": values["organic_orders"],
                    "ad_orders": values["ad_orders"], "ad_order_share": values["ad_order_share"],
                    "target_attainment": values["target_attainment"], "sessions": values["sessions"],
                    "organic_clicks": values["organic_clicks"], "impressions": values["impressions"],
                    "ad_clicks": values["ad_clicks"], "ctr": values["ctr"], "ad_cvr": values["ad_cvr"],
                    "spend": values["spend"], "ad_sales": values["ad_sales"],
                    "overall_cvr": values["overall_cvr"], "acos": values["acos"],
                }
                for field, value in metrics.items():
                    _write_metric(
                        worksheet, row_number, columns.get(field), value,
                        preserve=False,
                    )
                if values["sales"] is not None:
                    sales_units_total += values["sales"]
                    has_sales_units = True
                if values["_sales_amount"] is not None:
                    sales_amount_total += values["_sales_amount"]
                    has_sales_amount = True
                if values["spend"] is not None:
                    ad_spend_total += values["spend"]
                    has_spend = True
                daily_orders = (
                    reports["sales"].orders_by_sku_date.get(key, {}).get(item.date)
                    if reports["sales"].orders_by_sku_date is not None else None
                )
                daily_audit_entries.append(
                    {
                        "record_type": "daily_metrics",
                        "sku": sku,
                        "day_number": item.day_number,
                        "date": item.date.isoformat(),
                        "sales_source_row_present": sales_source_present,
                        "units_sold": _excel_number(values["sales"]),
                        "sales_amount": (
                            format(values["_sales_amount"], "f")
                            if api_mode and isinstance(values["_sales_amount"], Decimal)
                            else _excel_number(values["_sales_amount"])
                        ),
                        "total_order_count": (
                            format(daily_orders, "f")
                            if api_mode and isinstance(daily_orders, Decimal)
                            else _excel_number(daily_orders)
                        ),
                        "sales_source": "lingxing_openapi" if api_mode else "mapped_file",
                        "ad_spend_source_row_present": spend_source_present,
                        "actual_ad_spend": _excel_number(values["spend"]),
                    }
                )

            sku_warnings = list(info["warnings"])
            if not reports.get("actual_ad_spend") or not reports["actual_ad_spend"].rows_by_sku_date.get(key):
                sku_warnings.append("没有该 SKU 可确认的广告实际花费记录；花费单元格保持空白，不视为 0。")
            successful.append(
                SkuGenerationResult(
                    sku=sku,
                    status="partial" if api_mode else "success",
                    d01=d01.isoformat(),
                    d01_source=info["d01_source"],
                    latest_covered_date=latest.isoformat(),
                    days=len(output_timeline),
                    sales_units_total=float(sales_units_total) if has_sales_units else None,
                    sales_amount_total=(
                        format(sales_amount_total, "f") if api_mode and has_sales_amount
                        else (float(sales_amount_total) if has_sales_amount else None)
                    ),
                    actual_ad_spend_total=float(ad_spend_total) if has_spend else None,
                    warnings=sku_warnings,
                )
            )
        if not successful:
            raise ReportGenerationError("没有 SKU 成功生成；" + "; ".join(item.sku for item in failures))
        for name in source_sheets_to_remove:
            if name in workbook.sheetnames:
                workbook.remove(workbook[name])
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(stage_path)
        workbook.close()
        if output_path.exists():
            raise ReportGenerationError("输出文件在生成过程中出现，为保护已有周报，已停止覆盖。")
        os.replace(stage_path, output_path)
        results.extend(successful)
        results.extend(failures)
    except Exception:
        try:
            workbook.close()
        except (NameError, UnboundLocalError):
            pass
        stage_path.unlink(missing_ok=True)
        raise

    log_path = audit_target
    audit_records: list[dict[str, Any]] = [
        {
            "record_type": "run",
            "output_file": output_path.name,
            "template_file": template_path.name,
            "mapping_file": mapping_path.name if mapping_path is not None else None,
            "source_files": sorted({name for report in reports.values() for name in report.source_files}),
            "d01_mode": "manual_confirmed",
            "source_details": (
                {
                    "provider": "lingxing_openapi",
                    "store_name": api_store_name,
                    "store_sid": api_store_sid,
                    "queried_start": api_sales_windows[0].start_date.isoformat(),
                    "queried_end": api_sales_windows[-1].end_date.isoformat(),
                    "data_status": "sales_only_partial; ad_actual_spend_not_verified",
                    "d01_status": "manually_confirmed_by_user_for_store_sku_site_date",
                }
                if api_mode and api_sales_windows else None
            ),
            "report_notes": [note for report in reports.values() for note in (report.warnings or [])],
            "anomalies": [note for report in reports.values() for note in (report.anomalies or [])],
        }
    ]
    audit_records.extend(daily_audit_entries)
    actual = reports.get("actual_ad_spend")
    for item in results:
        audit_records.append({"record_type": "sku", **asdict(item)})
    if actual and actual.source_total_spend is not None:
        accepted_by_day = actual.accepted_spend_by_sku_date or {}
        for item in successful:
            sku_key = item.sku.casefold()
            if item.d01 and item.latest_covered_date:
                start = date.fromisoformat(item.d01)
                end = date.fromisoformat(item.latest_covered_date)
                source_window = sum(
                    (amount for day, amount in accepted_by_day.get(sku_key, {}).items() if start <= day <= end),
                    Decimal("0"),
                )
                has_source_rows = any(start <= day <= end for day in accepted_by_day.get(sku_key, {}))
                if has_source_rows:
                    generated_total = Decimal(str(item.actual_ad_spend_total)) if item.actual_ad_spend_total is not None else None
                    difference = generated_total - source_window if generated_total is not None else None
                    audit_records.append(
                        {
                            "record_type": "sku_ad_spend_reconciliation",
                            "sku": item.sku,
                            "d01": start.isoformat(),
                            "latest_covered_date": end.isoformat(),
                            "accepted_source_spend_in_window": float(source_window),
                            "generated_workbook_spend": float(generated_total) if generated_total is not None else None,
                            "difference": float(difference) if difference is not None else None,
                            "matches": difference == 0 if difference is not None else None,
                            "note": "差异表示存在数据排除或日期/店铺覆盖问题，需核对原始广告活动实际花费报表。",
                        }
                    )
        source_window_values = []
        for item in successful:
            if item.d01 and item.latest_covered_date:
                start = date.fromisoformat(item.d01)
                end = date.fromisoformat(item.latest_covered_date)
                values = accepted_by_day.get(item.sku.casefold(), {})
                if any(start <= day <= end for day in values):
                    source_window_values.append(sum((amount for day, amount in values.items() if start <= day <= end), Decimal("0")))
        selected_spend = sum(source_window_values, Decimal("0")) if source_window_values else None
        generated_values = [Decimal(str(item.actual_ad_spend_total)) for item in successful if item.actual_ad_spend_total is not None]
        generated_selected_spend = sum(generated_values, Decimal("0")) if generated_values else None
        requested_difference = (
            generated_selected_spend - selected_spend
            if generated_selected_spend is not None and selected_spend is not None else None
        )
        audit_records.append(
            {
                "record_type": "ad_spend_reconciliation",
                "source_total_of_parseable_rows": float(actual.source_total_spend),
                "accepted_source_total_after_duplicate_checks": float(actual.accepted_total_spend) if actual.accepted_total_spend is not None else None,
                "requested_sku_source_total_in_report_windows": float(selected_spend) if selected_spend is not None else None,
                "generated_requested_sku_total": float(generated_selected_spend) if generated_selected_spend is not None else None,
                "requested_sku_difference": float(requested_difference) if requested_difference is not None else None,
                "requested_sku_totals_match": requested_difference == 0 if requested_difference is not None else None,
                "all_source_skus_requested": bool(actual.all_skus) and actual.all_skus.issubset({sku.casefold() for sku in skus}),
                "note": "不同 SKU 子集时总额不可直接比较；请按原始报表核对活动实际花费汇总。",
            }
        )
    try:
        write_validation_log(log_path, audit_records)
    except Exception:
        # Keep the workbook, but make the missing audit artifact visible to the caller.
        raise ReportGenerationError(
            f"周报已生成，但数据校验日志无法写入：{log_path.name}。请不要重复覆盖，先检查输出目录。"
        )
    return GenerationResult(output_path, log_path, results)
