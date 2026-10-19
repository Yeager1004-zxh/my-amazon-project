from __future__ import annotations

import csv
import json
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any
from unittest.mock import patch
from urllib.parse import parse_qs

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Side

from amazon_weekly_report.ad_attribution import (
    AdCostEvidence,
    AdGroupCostEvidence,
    AttributionMapping,
    CampaignCostEvidence,
    PromotionTargetEvidence,
    attribute_exclusive_ad_group_costs,
    attribute_ad_costs,
)
from amazon_weekly_report.ad_preview import build_ad_spend_preview
from amazon_weekly_report.mcp_client import (
    BusinessTool,
    LingxingMcpClient,
    McpClientError,
    McpTool,
    VerifiedReadOnlyGrant,
    extract_business_tools,
    is_directory_search_tool,
    is_read_only_query_candidate,
    redact_sensitive,
    validate_json_schema,
)
from amazon_weekly_report.d01_settings import D01SettingsError, get_d01_dates, load_d01_settings, save_d01_dates
from amazon_weekly_report.providers.lingxing_api import (
    SB_AD_GROUP_REPORTS_PATH,
    SB_CAMPAIGN_REPORTS_PATH,
    SB_PRODUCT_ADS_PATH,
    SB_DIVIDE_ASIN_REPORTS_PATH,
    SD_CAMPAIGN_REPORTS_PATH,
    SD_PRODUCT_AD_REPORTS_PATH,
    SD_PRODUCT_ADS_PATH,
    SP_CAMPAIGN_REPORTS_PATH,
    SP_PRODUCT_AD_REPORTS_PATH,
    LingxingApiError,
    LingxingApiNotReady,
    LingxingOpenApiClient,
    AdProductCatalogRecord,
    ProductPerformanceMetrics,
    SbAdGroupSpend,
    SpAdSpendWindow,
    SpCampaignSpend,
    SpProductAdSpend,
    SkuDailyAdSpend,
    SkuDailySales,
    SkuDailySalesWindow,
    SkuProductMapping,
    parse_asin_daily_page,
    parse_product_performance_page,
    parse_product_performance_metrics_page,
    parse_sb_ad_group_report_page,
    parse_sb_campaign_report_page,
    parse_sb_product_ads_page,
    parse_sp_campaign_report_page,
    parse_sp_product_ad_report_page,
    parse_sd_campaign_report_page,
    parse_sd_product_ad_report_page,
    parse_sd_product_ads_page,
    parse_seller_list_response,
    split_date_windows,
)
from amazon_weekly_report.providers.lingxing_signing import (
    LingxingSignatureError,
    canonical_signing_string,
    encode_query_parameters,
    generate_signature,
    signed_post_query_parameters,
)
from amazon_weekly_report.report_generation import ReportGenerationError, generate_reports
from amazon_weekly_report.template_expansion import ensure_day_rows
from amazon_weekly_report.template_inspection import inspect_template
from amazon_weekly_report.timeline import build_timeline
from amazon_weekly_report.gui import WeeklyReportApp, _product_performance_value


HEADERS = [
    "周次", "天数", "周开始日期", "售价", "日销", "自然单", "广告单", "广告订单\n占比",
    "类目", "小类排名", "评分", "评论数", "日目标销量", "目标达成率", "Sessions",
    "自然点击", "广告曝光", "广告点击", "广告CTR", "广告CVR", "广告花费\n($)",
    "广告销售额\n($)", "整体CVR", "ACOS",
]


def create_template(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "新品推广周报"
    worksheet["A1"] = "推广阶段"
    worksheet["G1"] = "发货时间"
    worksheet.merge_cells("A2:A3")
    worksheet.merge_cells("G2:G3")
    worksheet["B2"] = "SAMPLE_ONLY"
    for column, value in enumerate(HEADERS, start=1):
        worksheet.cell(5, column, value)
    for day_number in range(1, 61):
        row = day_number + 5
        worksheet.cell(row, 1, f"第{(day_number - 1) // 7 + 1}周" if (day_number - 1) % 7 == 0 else None)
        worksheet.cell(row, 2, f"D{day_number:02d}")
        if (day_number - 1) % 7 == 0:
            worksheet.cell(row, 3, date(2026, 1, 1) + timedelta(days=day_number - 1))
        worksheet.cell(row, 5, 9876)  # recognizable sample value that must be cleared
        worksheet.cell(row, 13, None)
        worksheet.cell(row, 2).border = Border(bottom=Side(style="thin")) if day_number == 60 else Border()
    red_fill = "FFFF0000"
    worksheet.conditional_formatting.add("N6:N65", CellIsRule(operator="greaterThan", formula=["1"], fill=None))
    worksheet.conditional_formatting.add("X6:X65", CellIsRule(operator="greaterThan", formula=["1"], fill=None))
    worksheet.column_dimensions["B"].width = 12
    worksheet.row_dimensions[6].height = 23
    worksheet["C6"].number_format = "yyyymmdd"
    workbook.save(path)
    workbook.close()


class TemplateAndReportTests(unittest.TestCase):
    def test_report_generation_requires_a_manually_confirmed_d01_for_every_sku(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            create_template(template)
            data = root / "exports"
            data.mkdir()
            mapping = root / "mapping.json"
            mapping.write_text("{}", encoding="utf-8")
            with self.assertRaisesRegex(ReportGenerationError, "尚未设置并确认 D01"):
                generate_reports(
                    template_path=template,
                    data_folder=data,
                    mapping_path=mapping,
                    skus=["SKU_A", "SKU_B"],
                    output_path=root / "blocked.xlsx",
                    d01_confirmed=True,
                )
            self.assertFalse((root / "blocked.xlsx").exists())

    def test_existing_report_d01_change_requires_confirmation_then_rebases_dates_and_keeps_manual_fields(self) -> None:
        start = date(2026, 1, 1)
        mapping = SkuProductMapping("API-TEST", ("MSKU-A",), ())
        rows = tuple(
            SkuDailySales(987654, "API-TEST", "MSKU-A", "", start + timedelta(days=index), "USD", index + 1, index + 1, Decimal(f"{index + 1}.00"))
            for index in range(3)
        )
        window = SkuDailySalesWindow(start, start + timedelta(days=2), (mapping,), rows)
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            create_template(template)
            original = root / "original.xlsx"
            generate_reports(
                template_path=template,
                data_folder=None,
                mapping_path=None,
                skus=["API-TEST"],
                output_path=original,
                d01_dates={"API-TEST": start},
                d01_confirmed=True,
                api_sales_windows=[window],
                api_store_sid=987654,
                api_store_name="Synthetic Shop",
            )
            workbook = load_workbook(original)
            sheet = workbook["API-TEST"]
            sheet["A2"] = "保留推广阶段"
            sheet["G2"] = "保留发货时间"
            sheet["M6"] = 5
            sheet["M7"] = 6
            sheet["M8"] = 7
            workbook.save(original)
            workbook.close()

            new_d01 = start + timedelta(days=1)
            with self.assertRaisesRegex(ReportGenerationError, "日期错位风险"):
                generate_reports(
                    template_path=template,
                    data_folder=None,
                    mapping_path=None,
                    skus=["API-TEST"],
                    output_path=root / "unconfirmed.xlsx",
                    d01_dates={"API-TEST": new_d01},
                    d01_confirmed=True,
                    existing_report_path=original,
                    api_sales_windows=[window],
                    api_store_sid=987654,
                    api_store_name="Synthetic Shop",
                )

            updated = root / "confirmed.xlsx"
            generate_reports(
                template_path=template,
                data_folder=None,
                mapping_path=None,
                skus=["API-TEST"],
                output_path=updated,
                d01_dates={"API-TEST": new_d01},
                d01_confirmed=True,
                confirmed_d01_changes=True,
                existing_report_path=original,
                api_sales_windows=[window],
                api_store_sid=987654,
                api_store_name="Synthetic Shop",
            )
            workbook = load_workbook(updated)
            sheet = workbook["API-TEST"]
            self.assertEqual(sheet["C6"].value.date(), new_d01)
            self.assertEqual(sheet["B6"].value, "D01")
            self.assertEqual(sheet["B7"].value, "D02")
            self.assertEqual(sheet["B8"].value, "D03")
            self.assertEqual(sheet["E6"].value, 2)
            self.assertEqual(sheet["E7"].value, 3)
            self.assertIsNone(sheet["E8"].value)
            self.assertEqual(sheet["A2"].value, "保留推广阶段")
            self.assertEqual(sheet["G2"].value, "保留发货时间")
            self.assertEqual([sheet[f"M{row}"].value for row in (6, 7, 8)], [5, 6, 7])
            workbook.close()

    def test_gap_free_timeline_and_extension_past_sixty(self) -> None:
        first = date(2026, 1, 1)
        timeline = build_timeline(first, first + timedelta(days=65))
        self.assertEqual(len(timeline), 66)
        self.assertEqual([item.day_number for item in timeline], list(range(1, 67)))
        self.assertEqual(timeline[60].day_label, "D61")
        self.assertTrue(timeline[63].is_week_start)
        self.assertEqual(timeline[63].week_label, "第10周")

        with tempfile.TemporaryDirectory() as temp_dir:
            template = Path(temp_dir) / "template.xlsx"
            create_template(template)
            summary = inspect_template(template)
            workbook = load_workbook(template)
            worksheet = workbook[summary.sheet_name]
            rows = ensure_day_rows(worksheet, summary, 66)
            self.assertEqual(len(rows), 66)
            self.assertEqual(worksheet.cell(66, 2).style_id, worksheet.cell(10, 2).style_id)
            rows = ensure_day_rows(worksheet, summary, 70)
            self.assertEqual(len(rows), 70)
            ranges = {str(item.sqref) for item in worksheet.conditional_formatting}
            self.assertIn("N66:N71", ranges)
            self.assertIn("N72:N75", ranges)
            self.assertIn("X66:X71", ranges)
            workbook.close()

    def test_generation_clears_sample_and_keeps_missing_spend_blank(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            data = root / "exports"
            data.mkdir()
            create_template(template)
            sales_file = data / "sales.csv"
            with sales_file.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "销量", "销售额"])
                for index in range(10):
                    day = date(2026, 1, 1) + timedelta(days=index)
                    writer.writerow(["TEST_STORE", "TEST_SKU", day.isoformat(), 1, "10.00"])
            mapping = {
                "version": 1,
                "status": "SALES_ONLY_READY",
                "reports": {
                    "sales": {
                        "header_row": 1,
                        "sheet_name": "",
                        "file_patterns": ["sales.csv"],
                        "date_format": "%Y-%m-%d",
                        "source_headers": {
                            "store": "店铺", "sku": "SKU", "date": "日期",
                            "units_sold": "销量", "sales_amount": "销售额",
                        },
                    },
                    "actual_ad_spend": {
                        "header_row": None, "sheet_name": "", "sku_scope": "UNVERIFIED",
                        "file_patterns": [], "source_headers": {},
                    },
                    "ad_performance": {"enabled": False, "source_headers": {}},
                    "traffic": {"enabled": False, "source_headers": {}},
                    "product_snapshot": {"enabled": False, "source_headers": {}},
                },
            }
            mapping_path = root / "mapping.json"
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
            output = root / "report.xlsx"
            result = generate_reports(
                template_path=template,
                data_folder=data,
                mapping_path=mapping_path,
                skus=["TEST_SKU"],
                output_path=output,
                d01_dates={"TEST_SKU": date(2026, 1, 1)},
                d01_confirmed=True,
            )
            self.assertTrue(output.is_file())
            self.assertTrue(result.audit_path.is_file())
            self.assertEqual(result.sku_results[0].days, 10)
            workbook = load_workbook(output, data_only=False)
            worksheet = workbook["TEST_SKU"]
            self.assertEqual(worksheet["B2"].value, "TEST_SKU")
            self.assertIsNone(worksheet["A2"].value)
            self.assertIsNone(worksheet["G2"].value)
            self.assertIsNone(worksheet["B16"].value)
            self.assertIsNone(worksheet["E16"].value)
            self.assertNotEqual(worksheet["E6"].value, 9876)
            workbook.close()
            log_text = result.audit_path.read_text(encoding="utf-8")
            self.assertIn("2026-01-10", log_text)
            self.assertIn("广告花费单元格保持空白", log_text)
            daily_records = [json.loads(line) for line in log_text.splitlines()]
            first_day = next(item for item in daily_records if item.get("record_type") == "daily_metrics")
            self.assertEqual(first_day["sales_amount"], 10)
            self.assertTrue(first_day["sales_source_row_present"])
            self.assertIsNone(first_day["actual_ad_spend"])

            workbook = load_workbook(output)
            worksheet = workbook["TEST_SKU"]
            worksheet["A2"] = "D1-D14冷启动"
            worksheet["G2"] = "2026-01-05"
            worksheet["M6"] = 3
            worksheet["M16"] = 7
            workbook.save(output)
            workbook.close()

            with sales_file.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "销量", "销售额"])
                for index in range(61):
                    day = date(2026, 1, 1) + timedelta(days=index)
                    writer.writerow(["TEST_STORE", "TEST_SKU", day.isoformat(), 1, "10.00"])
            updated = root / "updated.xlsx"
            update_result = generate_reports(
                template_path=template,
                data_folder=data,
                mapping_path=mapping_path,
                skus=["TEST_SKU"],
                output_path=updated,
                d01_dates={"TEST_SKU": date(2026, 1, 1)},
                d01_confirmed=True,
                existing_report_path=output,
            )
            workbook = load_workbook(updated, data_only=False)
            worksheet = workbook["TEST_SKU"]
            self.assertEqual(worksheet["A2"].value, "D1-D14冷启动")
            self.assertEqual(worksheet["G2"].value, "2026-01-05")
            self.assertEqual(worksheet["M6"].value, 3)
            self.assertEqual(worksheet["M16"].value, 7)
            self.assertEqual(worksheet["N6"].value, 1 / 3)
            self.assertEqual(worksheet["B66"].value, "D61")
            self.assertEqual(worksheet["E66"].value, 1)
            self.assertIsNone(worksheet["U66"].value)
            self.assertIsNone(worksheet["M66"].value)
            self.assertIsNone(worksheet["E67"].value)
            workbook.close()
            self.assertIn("2026-03-02", update_result.audit_path.read_text(encoding="utf-8"))

    def test_actual_spend_sums_distinct_campaigns_and_flags_post_close_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            data = root / "exports"
            data.mkdir()
            create_template(template)
            with (data / "sales.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "销量", "销售额"])
                writer.writerow(["TEST_STORE", "TEST_SKU", "2026-01-01", 1, 10])
            with (data / "ads.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "活动ID", "实际花费", "关闭日期"])
                writer.writerow(["TEST_STORE", "TEST_SKU", "2026-01-02", "CAMPAIGN_A", "2.00", "2026-01-01"])
                writer.writerow(["TEST_STORE", "TEST_SKU", "2026-01-02", "CAMPAIGN_B", "3.00", "2026-01-01"])
            mapping = {
                "version": 1,
                "status": "READY_FOR_REVIEW",
                "reports": {
                    "sales": {
                        "header_row": 1, "sheet_name": "", "file_patterns": ["sales.csv"],
                        "date_format": "%Y-%m-%d",
                        "source_headers": {"store": "店铺", "sku": "SKU", "date": "日期", "units_sold": "销量", "sales_amount": "销售额"},
                    },
                    "actual_ad_spend": {
                        "header_row": 1, "sheet_name": "", "sku_scope": "SKU_LEVEL_CONFIRMED",
                        "coverage_scope": "ALL_AD_TYPES_TOTAL_CONFIRMED",
                        "file_patterns": ["ads.csv"], "date_format": "%Y-%m-%d",
                        "source_headers": {"store": "店铺", "sku": "SKU", "date": "日期", "actual_spend": "实际花费", "campaign_id": "活动ID", "campaign_disabled_at": "关闭日期"},
                    },
                    "ad_performance": {"enabled": False, "source_headers": {}},
                    "traffic": {"enabled": False, "source_headers": {}},
                    "product_snapshot": {"enabled": False, "source_headers": {}},
                },
            }
            mapping_path = root / "mapping.json"
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
            result = generate_reports(
                template_path=template,
                data_folder=data,
                mapping_path=mapping_path,
                skus=["TEST_SKU"],
                output_path=root / "report.xlsx",
                d01_dates={"TEST_SKU": date(2026, 1, 1)},
                d01_confirmed=True,
            )
            self.assertEqual(result.sku_results[0].status, "partial")
            self.assertIsNone(result.sku_results[0].actual_ad_spend_total)
            workbook = load_workbook(result.output_path, data_only=False)
            worksheet = workbook["TEST_SKU"]
            self.assertEqual(worksheet["U7"].value, 5)
            self.assertIsNone(worksheet["U6"].value)
            workbook.close()
            log_text = result.audit_path.read_text(encoding="utf-8")
            self.assertIn("关闭日期 2026-01-01 后仍有实际花费记录 2026-01-02", log_text)
            self.assertIn('"requested_sku_source_total_in_report_windows": 5.0', log_text)

    def test_repeated_campaign_cost_across_skus_is_left_unattributed(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            data = root / "exports"
            data.mkdir()
            create_template(template)
            with (data / "sales.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "销量", "销售额"])
                for sku in ("SKU_A", "SKU_B"):
                    writer.writerow(["TEST_STORE", sku, "2026-01-01", 1, 10])
            with (data / "ads.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "活动ID", "实际花费"])
                for sku in ("SKU_A", "SKU_B"):
                    writer.writerow(["TEST_STORE", sku, "2026-01-02", "SHARED_CAMPAIGN", "5.00"])
            mapping = {
                "version": 1,
                "status": "READY_FOR_REVIEW",
                "reports": {
                    "sales": {
                        "header_row": 1, "sheet_name": "", "file_patterns": ["sales.csv"],
                        "source_headers": {"store": "店铺", "sku": "SKU", "date": "日期", "units_sold": "销量", "sales_amount": "销售额"},
                    },
                    "actual_ad_spend": {
                        "header_row": 1, "sheet_name": "", "sku_scope": "SKU_LEVEL_CONFIRMED",
                        "coverage_scope": "ALL_AD_TYPES_TOTAL_CONFIRMED",
                        "file_patterns": ["ads.csv"],
                        "source_headers": {"store": "店铺", "sku": "SKU", "date": "日期", "actual_spend": "实际花费", "campaign_id": "活动ID"},
                    },
                    "ad_performance": {"enabled": False, "source_headers": {}},
                    "traffic": {"enabled": False, "source_headers": {}},
                    "product_snapshot": {"enabled": False, "source_headers": {}},
                },
            }
            mapping_path = root / "mapping.json"
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
            result = generate_reports(
                template_path=template,
                data_folder=data,
                mapping_path=mapping_path,
                skus=["SKU_A", "SKU_B"],
                output_path=root / "report.xlsx",
                d01_dates={"SKU_A": date(2026, 1, 1), "SKU_B": date(2026, 1, 1)},
                d01_confirmed=True,
            )
            workbook = load_workbook(result.output_path, data_only=False)
            for sku in ("SKU_A", "SKU_B"):
                self.assertIsNone(workbook[sku]["U7"].value)
            workbook.close()
            log_text = result.audit_path.read_text(encoding="utf-8")
            self.assertIn("相同活动/广告组花费重复关联多个 SKU", log_text)


class ApiReportGenerationTests(unittest.TestCase):
    def test_openapi_sales_generate_template_d61_and_keep_missing_distinct_from_zero(self) -> None:
        start = date(2026, 1, 1)
        end = start + timedelta(days=61)
        mapping = SkuProductMapping("API-TEST", ("MSKU-A", "MSKU-B"), ("ASINTEST0001",))
        rows: list[SkuDailySales] = []
        for offset in range((end - start).days + 1):
            event_date = start + timedelta(days=offset)
            for seller_sku in mapping.seller_skus:
                if offset == 0:
                    units, orders, amount = 0, 0, Decimal("0.00")
                elif offset == 1:
                    units, orders, amount = None, None, None
                else:
                    units, orders, amount = 1, 1, Decimal("10.00") if seller_sku == "MSKU-A" else Decimal("11.00")
                rows.append(SkuDailySales(
                    sid=987654,
                    local_sku="API-TEST",
                    seller_sku=seller_sku,
                    asin="ASINTEST0001",
                    event_date=event_date,
                    currency_code="USD",
                    units_sold=units,
                    order_count=orders,
                    sales_amount=amount,
                ))
        window = SkuDailySalesWindow(start, end, (mapping,), tuple(rows))

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            output = root / "api-report.xlsx"
            create_template(template)
            result = generate_reports(
                template_path=template,
                data_folder=None,
                mapping_path=None,
                skus=["API-TEST"],
                output_path=output,
                d01_dates={"API-TEST": start},
                d01_confirmed=True,
                api_sales_windows=[window],
                api_store_sid=987654,
                api_store_name="Synthetic Shop",
            )
            self.assertEqual(result.sku_results[0].status, "partial")
            self.assertEqual(result.sku_results[0].days, 62)
            self.assertTrue(any("未确认覆盖 SP、SB、SD" in warning for warning in result.sku_results[0].warnings))
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["API-TEST"]
            self.assertEqual(sheet["E6"].value, 0)
            self.assertIsNone(sheet["D6"].value)  # 0/0 is undefined, not zero price.
            self.assertIsNone(sheet["E7"].value)  # Missing API records stay blank.
            self.assertIsNone(sheet["D7"].value)
            self.assertEqual(sheet["E8"].value, 2)
            self.assertEqual(sheet["D8"].value, 10.5)
            self.assertEqual(sheet["B66"].value, "D61")
            self.assertEqual(sheet["E66"].value, 2)
            self.assertIsNone(sheet["U6"].value)  # Actual ad spend is not yet connected.
            self.assertIsNone(sheet["A2"].value)
            self.assertIsNone(sheet["G2"].value)
            self.assertIsNone(sheet["M6"].value)
            workbook.close()

            audit = [json.loads(line) for line in result.audit_path.read_text(encoding="utf-8").splitlines()]
            run = next(item for item in audit if item.get("record_type") == "run")
            self.assertEqual(run["source_details"]["provider"], "lingxing_openapi")
            self.assertEqual(run["source_details"]["store_sid"], 987654)
            self.assertTrue(any("广告" in note for note in run["report_notes"]))
            first_day = next(item for item in audit if item.get("record_type") == "daily_metrics" and item["day_number"] == 1)
            missing_day = next(item for item in audit if item.get("record_type") == "daily_metrics" and item["day_number"] == 2)
            self.assertEqual(first_day["total_order_count"], "0")
            self.assertEqual(first_day["sales_amount"], "0.00")
            self.assertIsNone(missing_day["total_order_count"])
            self.assertIsNone(missing_day["sales_amount"])
            self.assertEqual(first_day["sales_source"], "lingxing_openapi")

            workbook = load_workbook(output)
            sheet = workbook["API-TEST"]
            sheet["A2"] = "人工推广阶段"
            sheet["G2"] = "人工发货时间"
            sheet["M6"] = 4
            sheet["M66"] = 8
            workbook.save(output)
            workbook.close()

            added_date = end + timedelta(days=1)
            added_rows = tuple(
                SkuDailySales(987654, "API-TEST", seller_sku, "ASINTEST0001", added_date, "USD", 1, 1, Decimal("10.00"))
                for seller_sku in mapping.seller_skus
            )
            updated_window = SkuDailySalesWindow(start, added_date, (mapping,), (*rows, *added_rows))
            updated = root / "api-report-updated.xlsx"
            generate_reports(
                template_path=template,
                data_folder=None,
                mapping_path=None,
                skus=["API-TEST"],
                output_path=updated,
                d01_dates={"API-TEST": start},
                d01_confirmed=True,
                existing_report_path=output,
                api_sales_windows=[updated_window],
                api_store_sid=987654,
                api_store_name="Synthetic Shop",
            )
            workbook = load_workbook(updated)
            sheet = workbook["API-TEST"]
            self.assertEqual(sheet["A2"].value, "人工推广阶段")
            self.assertEqual(sheet["G2"].value, "人工发货时间")
            self.assertEqual(sheet["M6"].value, 4)
            self.assertEqual(sheet["M66"].value, 8)
            self.assertEqual(sheet["B68"].value, "D63")
            self.assertEqual(sheet["E68"].value, 2)
            workbook.close()

    def test_d01_requires_explicit_user_confirmation(self) -> None:
        start = date(2026, 1, 1)
        mapping = SkuProductMapping("API-TEST", ("MSKU-A",), ())
        row = SkuDailySales(987654, "API-TEST", "MSKU-A", "", start, "USD", 1, 1, Decimal("9.99"))
        window = SkuDailySalesWindow(start, start, (mapping,), (row,))
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            create_template(template)
            with self.assertRaisesRegex(ReportGenerationError, "尚未由用户明确确认"):
                generate_reports(
                    template_path=template,
                    data_folder=None,
                    mapping_path=None,
                    skus=["API-TEST"],
                    output_path=root / "blocked.xlsx",
                    d01_dates={"API-TEST": start},
                    api_sales_windows=[window],
                    api_store_sid=987654,
                    api_store_name="Synthetic Shop",
                )

    def test_partial_csv_ad_cost_is_kept_in_audit_but_never_written_as_total(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "template.xlsx"
            data = root / "exports"
            data.mkdir()
            create_template(template)
            with (data / "sales.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "销量", "销售额"])
                writer.writerow(["TEST_STORE", "TEST_SKU", "2026-02-01", 1, "10.00"])
                writer.writerow(["TEST_STORE", "TEST_SKU", "2026-02-02", 1, "11.00"])
            with (data / "ads.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["店铺", "SKU", "日期", "实际花费"])
                writer.writerow(["TEST_STORE", "TEST_SKU", "2026-02-02", "2.50"])
            mapping = {
                "version": 1,
                "status": "READY_FOR_REVIEW",
                "reports": {
                    "sales": {
                        "header_row": 1, "sheet_name": "", "file_patterns": ["sales.csv"],
                        "source_headers": {"store": "店铺", "sku": "SKU", "date": "日期", "units_sold": "销量", "sales_amount": "销售额"},
                    },
                    "actual_ad_spend": {
                        "header_row": 1, "sheet_name": "", "sku_scope": "SKU_LEVEL_CONFIRMED",
                        "coverage_scope": "PARTIAL_OR_UNKNOWN", "file_patterns": ["ads.csv"],
                        "source_headers": {"store": "店铺", "sku": "SKU", "date": "日期", "actual_spend": "实际花费"},
                    },
                    "ad_performance": {"enabled": False, "source_headers": {}},
                    "traffic": {"enabled": False, "source_headers": {}},
                    "product_snapshot": {"enabled": False, "source_headers": {}},
                },
            }
            mapping_path = root / "mapping.json"
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
            first = generate_reports(
                template_path=template, data_folder=data, mapping_path=mapping_path,
                skus=["TEST_SKU"], output_path=root / "first.xlsx",
                d01_dates={"TEST_SKU": date(2026, 2, 1)}, d01_confirmed=True,
            )
            self.assertEqual(first.sku_results[0].status, "partial")
            workbook = load_workbook(first.output_path)
            worksheet = workbook["TEST_SKU"]
            self.assertIsNone(worksheet["U7"].value)
            self.assertIsNone(worksheet["X7"].value)
            worksheet["U7"] = 99  # A previously generated total must not survive incomplete coverage.
            workbook.save(first.output_path)
            workbook.close()

            updated = generate_reports(
                template_path=template, data_folder=data, mapping_path=mapping_path,
                skus=["TEST_SKU"], output_path=root / "updated.xlsx",
                d01_dates={"TEST_SKU": date(2026, 2, 1)}, d01_confirmed=True,
                existing_report_path=first.output_path,
            )
            workbook = load_workbook(updated.output_path)
            self.assertIsNone(workbook["TEST_SKU"]["U7"].value)
            self.assertIsNone(workbook["TEST_SKU"]["X7"].value)
            workbook.close()
            daily = [json.loads(line) for line in updated.audit_path.read_text(encoding="utf-8").splitlines()]
            d02 = next(item for item in daily if item.get("record_type") == "daily_metrics" and item["day_number"] == 2)
            self.assertEqual(d02["ad_spend_source_value"], 2.5)
            self.assertEqual(d02["ad_spend_coverage_scope"], "PARTIAL_OR_UNKNOWN")
            self.assertIsNone(d02["actual_ad_spend"])
            sku = next(item for item in daily if item.get("record_type") == "sku")
            self.assertTrue(any("未确认覆盖 SP、SB、SD" in warning for warning in sku["warnings"]))


class D01SettingsTests(unittest.TestCase):
    def test_confirmed_d01_dates_persist_per_store_and_sku(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "d01_dates.json"
            first = date(2026, 8, 1)
            second = date(2026, 9, 3)
            save_d01_dates(510226, {"EX3106": first}, path=path)
            save_d01_dates(510226, {"ex3114": second}, path=path)
            save_d01_dates(510227, {"EX3106": second}, path=path)
            self.assertEqual(get_d01_dates(510226, ["ex3106", "EX3114"], path=path), {
                "ex3106": first, "ex3114": second,
            })
            self.assertEqual(get_d01_dates(510227, ["ex3106"], path=path), {"ex3106": second})
            self.assertEqual(set(load_d01_settings(path)), {"510226", "510227"})
            self.assertNotIn("secret", path.read_text(encoding="utf-8").casefold())

    def test_corrupt_d01_settings_fail_closed_without_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "d01_dates.json"
            path.write_text("{broken", encoding="utf-8")
            with self.assertRaises(D01SettingsError):
                save_d01_dates(510226, {"SKU_A": date(2026, 1, 1)}, path=path)
            self.assertEqual(path.read_text(encoding="utf-8"), "{broken")

    def test_d01_settings_reject_invalid_dates_and_sid(self) -> None:
        with self.assertRaises(D01SettingsError):
            save_d01_dates(True, {"SKU_A": date(2026, 1, 1)}, path=Path("unused.json"))
        with self.assertRaises(D01SettingsError):
            save_d01_dates(510226, {"SKU_A": "2026-01-01"}, path=Path("unused.json"))  # type: ignore[arg-type]


class SpAdApiTests(unittest.TestCase):
    @staticmethod
    def _envelope(records: list[dict[str, Any]], total: int | None = None) -> dict[str, Any]:
        return {
            "code": 0,
            "message": "success",
            "error_details": [],
            "request_id": "synthetic-request",
            "response_time": "2026-08-05 10:00:00",
            "total": len(records) if total is None else total,
            "data": records,
        }

    @staticmethod
    def _catalog_envelope(
        records: list[dict[str, Any]], total: int | None = None, next_token: str = ""
    ) -> dict[str, Any]:
        payload = SpAdApiTests._envelope(records, total)
        payload["next_token"] = next_token
        return payload

    def test_sp_report_parsers_validate_costs_ids_dates_and_business_code(self) -> None:
        report_day = date(2026, 8, 5)
        campaign = self._envelope([{
            "campaign_id": 123,
            "profile_id": 987,
            "cost": 3.25,
            "report_date": report_day.isoformat(),
        }])
        product = self._envelope([{
            "ad_id": 456,
            "asin": "ASINTEST0001",
            "sku": "MSKU-A",
            "ad_group_id": 789,
            "profile_id": 987,
            "campaign_id": 123,
            "cost": 3.25,
            "report_date": report_day.isoformat(),
        }])
        total, campaign_rows = parse_sp_campaign_report_page(campaign, report_date=report_day)
        self.assertEqual(total, 1)
        self.assertEqual(campaign_rows[0].cost, Decimal("3.25"))
        total, product_rows = parse_sp_product_ad_report_page(product, report_date=report_day)
        self.assertEqual(total, 1)
        self.assertEqual(product_rows[0].sku, "MSKU-A")
        self.assertEqual(product_rows[0].campaign_id, "123")

        wrong_code = dict(campaign, code=1)
        with self.assertRaisesRegex(LingxingApiError, "业务 code"):
            parse_sp_campaign_report_page(wrong_code, report_date=report_day)
        wrong_date = self._envelope([dict(campaign["data"][0], report_date="2026-08-06")])
        with self.assertRaisesRegex(LingxingApiError, "请求日期不一致"):
            parse_sp_campaign_report_page(wrong_date, report_date=report_day)
        wrong_cost = self._envelope([dict(product["data"][0], cost="3.25")])
        with self.assertRaisesRegex(LingxingApiError, "字段 cost 类型无效"):
            parse_sp_product_ad_report_page(wrong_cost, report_date=report_day)

    def test_sp_product_cost_is_unique_per_sku_and_cross_checked_to_campaign_total(self) -> None:
        report_day = date(2026, 8, 5)
        campaigns = [{"campaign_id": 100, "profile_id": 900, "cost": 3, "report_date": report_day.isoformat()}]
        products = [
            {"ad_id": 1, "asin": "ASINTEST0001", "sku": "MSKU-A", "ad_group_id": 10, "profile_id": 900, "campaign_id": 100, "cost": 1.25, "report_date": report_day.isoformat()},
            {"ad_id": 2, "asin": "ASINTEST0002", "sku": "MSKU-B", "ad_group_id": 10, "profile_id": 900, "campaign_id": 100, "cost": 1.75, "report_date": report_day.isoformat()},
        ]
        requests: list[dict[str, Any]] = []
        headers_seen: list[str | None] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            requests.append({"path": request.url.path, **body})
            headers_seen.append(request.headers.get("X-API-VERSION"))
            records = campaigns if request.url.path.endswith("spCampaignReports") else products
            return httpx.Response(200, json=self._envelope(records))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            result = client.query_sp_ad_spend_history(
                access_token="fixture-token",
                store_sid=987654,
                mappings=[
                    SkuProductMapping("LOCAL-A", ("MSKU-A",), ()),
                    SkuProductMapping("LOCAL-B", ("MSKU-B",), ()),
                ],
                start_date=report_day,
                end_date=report_day,
            )
        self.assertEqual(len(requests), 2)
        self.assertTrue(all(item["sid"] == 987654 and item["show_detail"] == 0 and item["length"] == 15 for item in requests))
        self.assertEqual(headers_seen, ["2", "2"])
        self.assertEqual(len(result[0].discrepancies), 0)
        self.assertEqual({item.local_sku: item.actual_spend for item in result[0].sku_daily_spends}, {
            "LOCAL-A": Decimal("1.25"), "LOCAL-B": Decimal("1.75"),
        })
        self.assertEqual(sum(item.actual_spend for item in result[0].sku_daily_spends), Decimal("3.00"))
        http.close()

    def test_sp_cost_mismatch_blocks_sku_attribution_and_pagination_is_complete(self) -> None:
        report_day = date(2026, 8, 5)
        campaign_records = [{"campaign_id": 100, "profile_id": 900, "cost": 4, "report_date": report_day.isoformat()}]
        product_records = [
            {"ad_id": index + 1, "asin": "ASINTEST0001", "sku": "MSKU-A", "ad_group_id": 10, "profile_id": 900, "campaign_id": 100, "cost": 0.25, "report_date": report_day.isoformat()}
            for index in range(16)
        ]
        request_offsets: list[tuple[str, int]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            request_offsets.append((request.url.path, body["offset"]))
            if request.url.path.endswith("spCampaignReports"):
                return httpx.Response(200, json=self._envelope(campaign_records))
            offset = body["offset"]
            return httpx.Response(200, json=self._envelope(product_records[offset:offset + 15], len(product_records)))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            result = client.query_sp_ad_spend_history(
                access_token="fixture-token",
                store_sid=987654,
                mappings=[SkuProductMapping("LOCAL-A", ("MSKU-A",), ())],
                start_date=report_day,
                end_date=report_day,
            )
        self.assertEqual(request_offsets, [
            ("/pb/openapi/newad/spCampaignReports", 0),
            ("/pb/openapi/newad/spProductAdReports", 0),
            ("/pb/openapi/newad/spProductAdReports", 15),
        ])
        self.assertEqual(len(result[0].product_rows), 16)
        self.assertEqual(len(result[0].discrepancies), 0)
        self.assertEqual(result[0].sku_daily_spends[0].actual_spend, Decimal("4.00"))
        http.close()

        def mismatch_handler(request: httpx.Request) -> httpx.Response:
            rows = campaign_records if request.url.path.endswith("spCampaignReports") else product_records[:1]
            return httpx.Response(200, json=self._envelope(rows))

        http = httpx.Client(transport=httpx.MockTransport(mismatch_handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            mismatch = client.query_sp_ad_spend_history(
                access_token="fixture-token",
                store_sid=987654,
                mappings=[SkuProductMapping("LOCAL-A", ("MSKU-A",), ())],
                start_date=report_day,
                end_date=report_day,
            )
        self.assertEqual(len(mismatch[0].discrepancies), 1)
        self.assertEqual(mismatch[0].sku_daily_spends, ())
        http.close()

    def test_sb_campaign_and_ad_group_reports_reconcile_without_sku_attribution(self) -> None:
        report_day = date(2026, 8, 28)
        campaign_rows = [{"campaign_id": 200, "profile_id": 900, "cost": 4.75, "report_date": report_day.isoformat()}]
        group_rows = [
            {"campaign_id": 200, "ad_group_id": 201, "profile_id": 900, "cost": 1.25, "report_date": report_day.isoformat()},
            {"campaign_id": 200, "ad_group_id": 202, "profile_id": 900, "cost": 3.50, "report_date": report_day.isoformat()},
        ]
        group_rows.extend(
            {"campaign_id": 200, "ad_group_id": 210 + index, "profile_id": 900, "cost": 0, "report_date": report_day.isoformat()}
            for index in range(14)
        )
        requests: list[tuple[str, dict[str, Any], str | None]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            requests.append((request.url.path, body, request.headers.get("X-API-VERSION")))
            if request.url.path == SB_PRODUCT_ADS_PATH:
                return httpx.Response(200, json=self._catalog_envelope([]))
            if request.url.path == SB_CAMPAIGN_REPORTS_PATH:
                return httpx.Response(200, json=self._envelope(campaign_rows))
            offset = body["offset"]
            return httpx.Response(200, json=self._envelope(group_rows[offset:offset + 15], len(group_rows)))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            windows = client.query_sb_ad_spend_history(
                access_token="fixture-token", store_sid=987654,
                start_date=report_day, end_date=report_day,
            )
        window = windows[0]
        self.assertEqual([item[0] for item in requests], [
            SB_PRODUCT_ADS_PATH, SB_CAMPAIGN_REPORTS_PATH, SB_AD_GROUP_REPORTS_PATH, SB_AD_GROUP_REPORTS_PATH,
        ])
        self.assertTrue(all(header == "2" for path, _body, header in requests if path != SB_PRODUCT_ADS_PATH))
        self.assertIsNone(requests[0][2])
        self.assertTrue(all(body["sid"] == 987654 and body["length"] == 15 for _path, body, _header in requests))
        self.assertEqual([body["offset"] for _path, body, _header in requests], [0, 0, 0, 15])
        self.assertTrue(all("show_detail" not in body for path, body, _header in requests if path != SB_PRODUCT_ADS_PATH))
        self.assertEqual(requests[3][1]["offset"], 15)
        self.assertEqual(window.ad_type, "SB")
        self.assertTrue(window.report_complete)
        self.assertEqual(window.campaign_cost_total, Decimal("4.75"))
        self.assertEqual(window.detail_cost_total, Decimal("4.75"))
        self.assertEqual(window.unattributed_cost, Decimal("4.75"))
        self.assertEqual(len(window.ad_group_rows), 16)
        self.assertEqual(window.sku_daily_spends, ())
        http.close()

    def test_current_sb_catalog_is_candidate_only_for_historical_report_date(self) -> None:
        report_day = date(2026, 8, 29)
        requests: list[str] = []
        campaign = [{"campaign_id": 10, "profile_id": 20, "cost": 19.21, "report_date": report_day.isoformat()}]
        groups = [{"campaign_id": 10, "ad_group_id": 11, "profile_id": 20, "cost": 19.21, "report_date": report_day.isoformat()}]
        catalog = [{
            "campaign_id": 10, "ad_group_id": 11, "ad_creative_id": 12,
            "profile_id": 20, "state": "ENABLED", "serving_status": "AD_STATUS_LIVE",
            "creation_date": 1704643200000, "last_updated_date": 1704643200000,
            "asin": ["ASIN-A"], "creative_type": "PRODUCT_COLLECTION",
        }]

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request.url.path)
            if request.url.path == SB_PRODUCT_ADS_PATH:
                return httpx.Response(200, json=self._catalog_envelope(catalog))
            if request.url.path == SB_CAMPAIGN_REPORTS_PATH:
                return httpx.Response(200, json=self._envelope(campaign))
            return httpx.Response(200, json=self._envelope(groups))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            windows = client.query_sb_ad_spend_history(
                access_token="fixture-token", store_sid=77,
                start_date=report_day, end_date=report_day,
            )
        window = windows[0]
        self.assertEqual(requests, [SB_PRODUCT_ADS_PATH, SB_CAMPAIGN_REPORTS_PATH, SB_AD_GROUP_REPORTS_PATH])
        self.assertTrue(window.report_complete)
        self.assertTrue(window.promotion_catalog_complete)
        self.assertEqual(window.promotion_targets[0].asin, "ASIN-A")
        self.assertFalse(window.promotion_targets[0].identity_valid_for_report)
        summaries, details = build_ad_spend_preview(
            windows_by_type={"SB": windows},
            mappings=[SkuProductMapping("LOCAL-A", (), ("ASIN-A",))],
            store_sid=77, requested_skus=["LOCAL-A"], start_date=report_day, end_date=report_day,
        )
        sku_row = next(row for row in summaries if row.local_sku == "LOCAL-A")
        catalog_row = next(row for row in details if row.report_level.startswith("当前商品目录"))
        self.assertEqual(sku_row.sb_cost, "")
        self.assertIn("历史", sku_row.integrity)
        self.assertEqual(catalog_row.local_sku, "LOCAL-A")
        self.assertEqual(catalog_row.cost, "")
        self.assertIn("不参与费用归属", catalog_row.integrity)
        http.close()

    def test_sd_product_cost_is_mapped_only_after_campaign_reconciliation(self) -> None:
        report_day = date(2026, 8, 29)
        campaign_rows = [{"campaign_id": 300, "profile_id": 900, "cost": 3, "report_date": report_day.isoformat()}]
        product_rows = [
            {"ad_id": 301, "asin": "ASINTEST0001", "sku": "MSKU-A", "ad_group_id": 302, "profile_id": 900, "campaign_id": 300, "cost": 2.25, "report_date": report_day.isoformat()},
            {"ad_id": 303, "asin": "ASINTEST0002", "sku": "UNMAPPED-MSKU", "ad_group_id": 304, "profile_id": 900, "campaign_id": 300, "cost": 0.75, "report_date": report_day.isoformat()},
        ]
        requests: list[tuple[str, dict[str, Any], str | None]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            requests.append((request.url.path, body, request.headers.get("X-API-VERSION")))
            if request.url.path == SD_PRODUCT_ADS_PATH:
                return httpx.Response(200, json=self._catalog_envelope([{
                    "ad_id": 399, "ad_group_id": 398, "campaign_id": 397, "state": "enabled",
                    "sku": "MSKU-A", "asin": "ASINTEST0001", "profile_id": 900,
                    "creation_date": 1640591614, "last_updated_date": 1640591614,
                    "serving_status": "AD_STATUS_LIVE",
                }]))
            records = campaign_rows if request.url.path == SD_CAMPAIGN_REPORTS_PATH else product_rows
            return httpx.Response(200, json=self._envelope(records))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            windows = client.query_sd_ad_spend_history(
                access_token="fixture-token", store_sid=987654,
                mappings=[SkuProductMapping("LOCAL-A", ("MSKU-A",), ("ASINTEST0001",))],
                start_date=report_day, end_date=report_day,
            )
        window = windows[0]
        self.assertEqual([item[0] for item in requests], [SD_CAMPAIGN_REPORTS_PATH, SD_PRODUCT_AD_REPORTS_PATH, SD_PRODUCT_ADS_PATH])
        self.assertTrue(all(header == "2" for _path, _body, header in requests))
        self.assertTrue(all(body["show_detail"] == 0 and body["offset"] == 0 for path, body, _header in requests if path != SD_PRODUCT_ADS_PATH))
        self.assertTrue(window.report_complete)
        self.assertEqual(window.product_catalog_rows[0].ad_id, "399")
        self.assertEqual(window.campaign_cost_total, Decimal("3"))
        self.assertEqual(window.detail_cost_total, Decimal("3.00"))
        self.assertEqual(window.sku_daily_spends[0].actual_spend, Decimal("2.25"))
        self.assertEqual(window.unattributed_cost, Decimal("0.75"))
        http.close()

    def test_sd_historical_product_rows_attribute_without_catalog_and_block_identity_conflicts(self) -> None:
        first_day = date(2026, 10, 10)
        second_day = date(2026, 10, 11)
        campaign_by_date = {
            first_day: [{"campaign_id": 300, "profile_id": 900, "cost": 3, "report_date": first_day.isoformat()}],
            second_day: [{"campaign_id": 301, "profile_id": 900, "cost": 4, "report_date": second_day.isoformat()}],
        }
        products_by_date = {
            first_day: [
                {"ad_id": 401, "asin": "ASIN-A", "sku": "MSKU-A", "ad_group_id": 402, "profile_id": 900, "campaign_id": 300, "cost": 2, "report_date": first_day.isoformat()},
                {"ad_id": 403, "asin": "ASIN-B", "sku": "MSKU-A", "ad_group_id": 404, "profile_id": 900, "campaign_id": 300, "cost": 1, "report_date": first_day.isoformat()},
            ],
            second_day: [
                {"ad_id": 405, "asin": "ASIN-A", "sku": "MSKU-A", "ad_group_id": 406, "profile_id": 900, "campaign_id": 301, "cost": 4, "report_date": second_day.isoformat()},
            ],
        }
        request_paths: list[str] = []

        def handler(request: httpx.Request) -> httpx.Response:
            request_paths.append(request.url.path)
            body = json.loads(request.content)
            if request.url.path == SD_PRODUCT_ADS_PATH:
                return httpx.Response(503, json={"message": "catalog unavailable"})
            day = date.fromisoformat(body["report_date"])
            rows = campaign_by_date[day] if request.url.path == SD_CAMPAIGN_REPORTS_PATH else products_by_date[day]
            return httpx.Response(200, json=self._envelope(rows))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        mappings = [
            SkuProductMapping("LOCAL-A", ("MSKU-A",), ("ASIN-A",)),
            SkuProductMapping("LOCAL-B", ("MSKU-B",), ("ASIN-B",)),
        ]
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            windows = client.query_sd_ad_spend_history(
                access_token="fixture-token", store_sid=987654, mappings=mappings,
                start_date=first_day, end_date=second_day,
            )
        self.assertEqual(len(windows), 2)
        self.assertEqual(request_paths, [
            SD_CAMPAIGN_REPORTS_PATH, SD_PRODUCT_AD_REPORTS_PATH,
            SD_CAMPAIGN_REPORTS_PATH, SD_PRODUCT_AD_REPORTS_PATH,
            SD_PRODUCT_ADS_PATH,
        ])
        self.assertTrue(all(window.report_complete for window in windows))
        self.assertEqual(windows[0].sku_daily_spends[0].local_sku, "LOCAL-A")
        self.assertEqual(windows[0].sku_daily_spends[0].actual_spend, Decimal("2"))
        self.assertEqual(windows[0].unattributed_cost, Decimal("1"))
        self.assertEqual(windows[1].sku_daily_spends[0].report_date, second_day)
        self.assertEqual(windows[1].sku_daily_spends[0].actual_spend, Decimal("4"))
        self.assertEqual(windows[0].product_catalog_rows, [])
        http.close()

    def test_sp_campaign_matching_uses_profile_and_campaign_ids_together(self) -> None:
        report_day = date(2026, 8, 29)
        campaign_rows = [
            {"campaign_id": 500, "profile_id": 900, "cost": 1.25, "report_date": report_day.isoformat()},
            {"campaign_id": 500, "profile_id": 901, "cost": 2.75, "report_date": report_day.isoformat()},
        ]
        product_rows = [
            {"ad_id": 700, "asin": "ASINTEST0001", "sku": "MSKU-A", "ad_group_id": 701, "profile_id": 900, "campaign_id": 500, "cost": 1.25, "report_date": report_day.isoformat()},
            {"ad_id": 700, "asin": "ASINTEST0002", "sku": "MSKU-B", "ad_group_id": 702, "profile_id": 901, "campaign_id": 500, "cost": 2.75, "report_date": report_day.isoformat()},
        ]

        def handler(request: httpx.Request) -> httpx.Response:
            records = campaign_rows if request.url.path == SP_CAMPAIGN_REPORTS_PATH else product_rows
            return httpx.Response(200, json=self._envelope(records))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            window = client.query_sp_ad_spend_history(
                access_token="fixture-token", store_sid=987654,
                mappings=[
                    SkuProductMapping("LOCAL-A", ("MSKU-A",), ()),
                    SkuProductMapping("LOCAL-B", ("MSKU-B",), ()),
                ],
                start_date=report_day, end_date=report_day,
            )[0]
        self.assertTrue(window.report_complete)
        self.assertEqual(len(window.campaign_rows), 2)
        self.assertEqual({row.profile_id for row in window.campaign_rows}, {"900", "901"})
        self.assertEqual({row.local_sku: row.actual_spend for row in window.sku_daily_spends}, {
            "LOCAL-A": Decimal("1.25"), "LOCAL-B": Decimal("2.75"),
        })
        self.assertEqual(window.campaign_cost_total, Decimal("4.00"))
        http.close()

    def test_zero_cost_record_is_distinct_from_empty_report_data(self) -> None:
        report_day = date(2026, 8, 28)

        def query_with_rows(campaign_rows: list[dict[str, Any]], product_rows: list[dict[str, Any]]):
            def handler(request: httpx.Request) -> httpx.Response:
                records = campaign_rows if request.url.path == SP_CAMPAIGN_REPORTS_PATH else product_rows
                return httpx.Response(200, json=self._envelope(records))

            http = httpx.Client(transport=httpx.MockTransport(handler))
            with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
                return client.query_sp_ad_spend_history(
                    access_token="fixture-token", store_sid=987654,
                    mappings=[SkuProductMapping("LOCAL-A", ("MSKU-A",), ())],
                    start_date=report_day, end_date=report_day,
                )[0]

        campaign = [{"campaign_id": 400, "profile_id": 900, "cost": 0, "report_date": report_day.isoformat()}]
        product = [{"ad_id": 401, "asin": "ASINTEST0001", "sku": "MSKU-A", "ad_group_id": 402, "profile_id": 900, "campaign_id": 400, "cost": 0, "report_date": report_day.isoformat()}]
        explicit_zero = query_with_rows(campaign, product)
        self.assertTrue(explicit_zero.report_complete)
        self.assertEqual(explicit_zero.sku_daily_spends[0].actual_spend, Decimal("0"))
        self.assertEqual(explicit_zero.unattributed_cost, Decimal("0"))

        no_rows = query_with_rows([], [])
        self.assertFalse(no_rows.report_complete)
        self.assertIsNone(no_rows.campaign_cost_total)
        self.assertIsNone(no_rows.detail_cost_total)
        self.assertIsNone(no_rows.unattributed_cost)
        self.assertEqual(no_rows.sku_daily_spends, ())

    def test_sb_sd_parsers_validate_identifiers_costs_and_dates(self) -> None:
        report_day = date(2026, 8, 28)
        campaign = self._envelope([{"campaign_id": 10, "profile_id": 20, "cost": 1.125, "report_date": report_day.isoformat()}])
        group = self._envelope([{"campaign_id": 10, "ad_group_id": 11, "profile_id": 20, "cost": 1.125, "report_date": report_day.isoformat()}])
        product = self._envelope([{"ad_id": 12, "asin": "ASINTEST0001", "sku": "MSKU-A", "ad_group_id": 11, "profile_id": 20, "campaign_id": 10, "cost": 1.125, "report_date": report_day.isoformat()}])
        self.assertEqual(parse_sb_campaign_report_page(campaign, report_date=report_day)[1][0].cost, Decimal("1.125"))
        self.assertEqual(parse_sb_ad_group_report_page(group, report_date=report_day)[1][0].ad_group_id, "11")
        self.assertEqual(parse_sd_campaign_report_page(campaign, report_date=report_day)[1][0].campaign_id, "10")
        self.assertEqual(parse_sd_product_ad_report_page(product, report_date=report_day)[1][0].sku, "MSKU-A")
        with self.assertRaises(LingxingApiError):
            parse_sb_ad_group_report_page(self._envelope([dict(group["data"][0], cost="1.125")]), report_date=report_day)

    def test_sb_sd_product_catalog_parsers_validate_real_schema_fields(self) -> None:
        sb = self._catalog_envelope([{
            "campaign_id": 101, "ad_group_id": 102, "ad_creative_id": 103,
            "creative_id": 9, "name": "synthetic", "state": "ENABLED",
            "creation_date": 1704643200000, "last_updated_date": 1704643200000,
            "profile_id": 104, "serving_status": "AD_STATUS_LIVE",
            "asin": ["ASIN-A", "ASIN-B"], "creative_type": "PRODUCT_COLLECTION",
            "creative": {"asins": ["IGNORED-DETAIL-ASIN"]},
            "landing_page": {"asins": ["IGNORED-LANDING-ASIN"]},
            "video_creatives": [{"videos": [{"versions": [{"url": "https://example.invalid/ignored.mp4"}]}]}],
        }])
        total, sb_rows, token = parse_sb_product_ads_page(sb)
        self.assertEqual((total, token), (1, ""))
        self.assertEqual(sb_rows[0].asins, ("ASIN-A", "ASIN-B"))
        self.assertEqual(sb_rows[0].ad_creative_id, "103")
        self.assertEqual(sb_rows[0].creation_date_value, Decimal("1704643200000"))

        sd = self._catalog_envelope([{
            "ad_id": 201, "ad_group_id": 202, "campaign_id": 203, "state": "enabled",
            "sku": "MSKU-A", "asin": "ASIN-C", "profile_id": 204,
            "creation_date": 1640591614, "last_updated_date": 1640591614,
            "serving_status": "AD_STATUS_LIVE",
        }])
        total, sd_rows, token = parse_sd_product_ads_page(sd)
        self.assertEqual((total, token, sd_rows[0].seller_sku, sd_rows[0].asins), (1, "", "MSKU-A", ("ASIN-C",)))

        with self.assertRaises(LingxingApiError):
            parse_sb_product_ads_page(self._catalog_envelope([dict(sb["data"][0], asin="ASIN-A")]))
        with self.assertRaises(LingxingApiError):
            parse_sd_product_ads_page(self._catalog_envelope([dict(sd["data"][0], ad_id="not-numeric")]))
        with self.assertRaises(LingxingApiError):
            parse_sd_product_ads_page(dict(sd, next_token=None))

    def test_sb_product_catalog_cursor_pagination_is_read_only_and_complete(self) -> None:
        records = [{
            "campaign_id": 100 + index, "ad_group_id": 200 + index,
            "ad_creative_id": 300 + index, "state": "ENABLED",
            "creation_date": 1704643200000, "last_updated_date": 1704643200000,
            "profile_id": 900, "serving_status": "AD_STATUS_LIVE",
            "asin": [f"ASIN-{index}"], "creative_type": "PRODUCT_COLLECTION",
        } for index in range(16)]
        requests: list[tuple[dict[str, Any], str | None]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content)
            requests.append((body, request.headers.get("X-API-VERSION")))
            cursor_request = body.get("next_token") == "cursor-1"
            rows = records[15:] if cursor_request else records[:15]
            next_token = "" if cursor_request else "cursor-1"
            return httpx.Response(200, json=self._catalog_envelope(rows, 16, next_token))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            rows = client.query_sb_product_ads(access_token="fixture-token", store_sid=987654)
        self.assertEqual(len(rows), 16)
        self.assertEqual(requests[0][0], {"sid": 987654, "offset": 0, "length": 15})
        self.assertEqual(requests[1][0], {"sid": 987654, "length": 15, "next_token": "cursor-1"})
        self.assertEqual([header for _body, header in requests], [None, None])
        http.close()

    def test_malformed_sb_divide_contract_is_blocked_before_network(self) -> None:
        requests: list[httpx.Request] = []
        http = httpx.Client(transport=httpx.MockTransport(
            lambda request: requests.append(request) or httpx.Response(200, json={})
        ))
        with LingxingOpenApiClient("1234567890123456", "fixture-app-secret", http_client=http) as client:
            with self.assertRaises(LingxingApiNotReady):
                client.query_sb_divide_asin_reports(access_token="fixture-token", profile_id=900, report_date="2026-08-29")
        self.assertEqual(requests, [])
        http.close()


class AdPreviewTests(unittest.TestCase):
    def test_summary_shows_sku_candidates_and_type_separated_unallocated_costs(self) -> None:
        report_day = date(2026, 8, 28)
        mapping = SkuProductMapping("LOCAL-A", ("MSKU-A",), ())
        sp = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "1", Decimal("1.50")),),
            (
                SpProductAdSpend(report_day, "MSKU-A", "ASINTEST0001", "11", "12", "1", Decimal("1.00")),
                SpProductAdSpend(report_day, "UNKNOWN", "ASINTEST0002", "13", "14", "1", Decimal("0.50")),
            ),
            (SkuDailyAdSpend(1, "LOCAL-A", "MSKU-A", report_day, Decimal("1.00"), 1),),
            (), ad_type="SP", campaign_cost_total=Decimal("1.50"),
            detail_cost_total=Decimal("1.50"), unattributed_cost=Decimal("0.50"), report_complete=True,
        )
        sb = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "2", Decimal("4.00")),), (), (), (),
            ad_type="SB", ad_group_rows=(SbAdGroupSpend(report_day, "2", "21", Decimal("4.00")),),
            campaign_cost_total=Decimal("4.00"), detail_cost_total=Decimal("4.00"),
            unattributed_cost=Decimal("4.00"), report_complete=True,
        )
        sd = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "3", Decimal("2.25")),),
            (
                SpProductAdSpend(report_day, "MSKU-A", "ASINTEST0001", "31", "32", "3", Decimal("2.00")),
                SpProductAdSpend(report_day, "UNKNOWN-SD", "ASINTEST0002", "33", "34", "3", Decimal("0.25")),
            ),
            (SkuDailyAdSpend(1, "LOCAL-A", "MSKU-A", report_day, Decimal("2.00"), 1),),
            (), ad_type="SD", campaign_cost_total=Decimal("2.25"),
            detail_cost_total=Decimal("2.25"), unattributed_cost=Decimal("0.25"), report_complete=True,
        )
        summaries, details = build_ad_spend_preview(
            windows_by_type={"SP": [sp], "SB": [sb], "SD": [sd]},
            mappings=[mapping], requested_skus=["LOCAL-A"],
            store_sid=1,
            start_date=report_day, end_date=report_day,
        )
        sku_row, unallocated_row = summaries
        self.assertEqual((sku_row.sp_cost, sku_row.sb_cost, sku_row.sd_cost), ("1.00", "", "2.00"))
        self.assertEqual(sku_row.confirmed_total, "")
        self.assertEqual(sku_row.unattributed_cost, "")
        self.assertIn("SB 活动/广告组成本已核平；创意目录未确认", sku_row.integrity)
        self.assertIn("币种/日期口径/历史映射时效尚未全部验收", sku_row.integrity)
        self.assertEqual((unallocated_row.sp_cost, unallocated_row.sb_cost, unallocated_row.sd_cost), ("0.50", "4.00", "0.25"))
        self.assertEqual(unallocated_row.confirmed_total, "")
        self.assertIn("未归属费用按类型分开显示", unallocated_row.integrity)
        self.assertTrue(any(row.report_level == "活动总额（仅核对）" and "不与明细相加" in row.integrity for row in details))
        self.assertTrue(any(
            row.report_level == "广告组记录（不重复计费）" and not row.local_sku and row.cost == "4.00"
            for row in details
        ))
        self.assertTrue(any(row.report_sku == "UNKNOWN" and row.local_sku == "" for row in details))

    def test_preview_uses_exact_asin_when_report_sku_is_missing_and_keeps_id_evidence(self) -> None:
        report_day = date(2026, 8, 29)
        sb = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "campaign-1", Decimal("19.21"), "profile-1"),),
            (SpProductAdSpend(report_day, "", "ASIN-A", "ad-1", "group-1", "campaign-1", Decimal("19.21"), "profile-1"),),
            (), (), ad_type="SB", campaign_cost_total=Decimal("19.21"),
            detail_cost_total=Decimal("19.21"), unattributed_cost=Decimal("0"), report_complete=True,
        )
        summaries, details = build_ad_spend_preview(
            windows_by_type={"SB": [sb]},
            mappings=[SkuProductMapping("LOCAL-A", (), ("ASIN-A",))],
            store_sid=77,
            requested_skus=["LOCAL-A"],
            start_date=report_day,
            end_date=report_day,
        )
        sku_row = next(row for row in summaries if row.local_sku == "LOCAL-A")
        sb_detail = next(row for row in details if row.ad_type == "SB" and row.report_level == "商品/分摊明细")
        self.assertEqual(sku_row.sb_cost, "19.21")
        self.assertEqual((sb_detail.local_sku, sb_detail.asin, sb_detail.ad_id), ("LOCAL-A", "ASIN-A", "ad-1"))
        self.assertIn("ASIN", sb_detail.integrity)

    def test_preview_automatically_attributes_exclusive_sb_ad_group(self) -> None:
        report_day = date(2026, 8, 29)
        sb = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "campaign-1", Decimal("19.21"), "profile-1"),),
            (), (), (), ad_type="SB",
            ad_group_rows=(SbAdGroupSpend(report_day, "campaign-1", "group-1", Decimal("19.21"), "profile-1"),),
            campaign_cost_total=Decimal("19.21"), detail_cost_total=Decimal("19.21"),
            unattributed_cost=Decimal("19.21"), report_complete=True,
            promotion_targets=(
                PromotionTargetEvidence("SB", 77, report_day, "profile-1", "campaign-1", "group-1", "ad-1", "", "ASIN-A", identity_valid_for_report=True),
            ),
            promotion_catalog_complete=True,
        )
        summaries, details = build_ad_spend_preview(
            windows_by_type={"SB": [sb]}, mappings=[SkuProductMapping("LOCAL-A", (), ("ASIN-A",))],
            store_sid=77, requested_skus=["LOCAL-A"], start_date=report_day, end_date=report_day,
        )
        sku_row = next(row for row in summaries if row.local_sku == "LOCAL-A")
        group_row = next(row for row in details if row.report_level == "广告组费用")
        self.assertEqual(sku_row.sb_cost, "19.21")
        self.assertEqual(group_row.local_sku, "LOCAL-A")

    def test_missing_report_is_not_zero_and_explicit_sku_zero_is_displayed(self) -> None:
        report_day = date(2026, 8, 29)
        zero = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "1", Decimal("0")),),
            (SpProductAdSpend(report_day, "MSKU-A", "ASINTEST0001", "11", "12", "1", Decimal("0")),),
            (SkuDailyAdSpend(1, "LOCAL-A", "MSKU-A", report_day, Decimal("0"), 1),),
            (), ad_type="SP", campaign_cost_total=Decimal("0"), detail_cost_total=Decimal("0"),
            unattributed_cost=Decimal("0"), report_complete=True,
        )
        summaries, _details = build_ad_spend_preview(
            windows_by_type={"SP": [zero]}, mappings=[SkuProductMapping("LOCAL-A", ("MSKU-A",), ())],
            store_sid=1,
            requested_skus=["LOCAL-A"], start_date=report_day, end_date=report_day,
        )
        sku_row, unallocated_row = summaries
        self.assertEqual(sku_row.sp_cost, "0")
        self.assertIn("SB 查询结果缺失（未知）", sku_row.integrity)
        self.assertIn("SD 查询结果缺失（未知）", sku_row.integrity)
        self.assertIn("SB=查询结果缺失", unallocated_row.unattributed_cost)
        self.assertEqual(unallocated_row.sb_cost, "")

    def test_sd_activity_without_historical_product_rows_remains_unknown(self) -> None:
        report_day = date(2026, 8, 29)
        sd = SpAdSpendWindow(
            report_day, report_day,
            (SpCampaignSpend(report_day, "campaign-1", Decimal("4.50"), "profile-1"),),
            (), (), (), ad_type="SD", campaign_cost_total=Decimal("4.50"),
            detail_cost_total=None, unattributed_cost=None, report_complete=False,
            product_catalog_rows=(AdProductCatalogRecord(
                "SD", "profile-1", "campaign-1", "group-1", "ad-1",
                seller_sku="MSKU-A", asins=("ASIN-A",),
            ),),
        )
        summaries, _details = build_ad_spend_preview(
            windows_by_type={"SD": [sd]}, mappings=[SkuProductMapping("LOCAL-A", ("MSKU-A",), ("ASIN-A",))],
            store_sid=77, requested_skus=["LOCAL-A"], start_date=report_day, end_date=report_day,
        )
        sku_row = next(row for row in summaries if row.local_sku == "LOCAL-A")
        self.assertEqual(sku_row.sd_cost, "")
        self.assertIn("SD 活动报表有 1 条记录，商品报表无记录", sku_row.integrity)
        self.assertIn("SKU 花费未知", sku_row.integrity)


class AdAttributionTests(unittest.TestCase):
    def _campaign(self, campaign_id: str, cost: str, *, profile: str = "profile-1") -> CampaignCostEvidence:
        return CampaignCostEvidence("SP", 10, date(2026, 8, 29), profile, campaign_id, Decimal(cost))

    def _line(
        self,
        campaign_id: str,
        cost: str,
        *,
        seller_sku: str = "",
        asin: str = "",
        ad_id: str = "ad-1",
        profile: str = "profile-1",
    ) -> AdCostEvidence:
        return AdCostEvidence(
            "SP", 10, date(2026, 8, 29), profile, campaign_id, Decimal(cost),
            "group-1", ad_id, seller_sku, asin,
        )

    def test_exact_msku_and_asin_maps_multiple_aliases_to_one_sku(self) -> None:
        result = attribute_ad_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", ("MSKU-OLD", "MSKU-NEW"), ("ASIN-OLD", "ASIN-NEW"))],
            product_costs=[
                self._line("campaign-1", "7.21", seller_sku="msku-old", asin="ASIN-OLD", ad_id="ad-1"),
                self._line("campaign-1", "12.00", seller_sku="MSKU-NEW", asin="asin-new", ad_id="ad-2"),
            ],
            campaign_costs=[self._campaign("campaign-1", "19.21")],
        )
        self.assertEqual([row.local_sku for row in result.rows], ["LOCAL-A", "LOCAL-A"])
        self.assertEqual([row.matched_by for row in result.rows], ["MSKU+ASIN", "MSKU+ASIN"])
        self.assertEqual(result.daily_costs()[("SP", date(2026, 8, 29), "local-a")], Decimal("19.21"))
        self.assertFalse(any(row.formal_ready for row in result.rows))
        self.assertIn(("SP", 10, date(2026, 8, 29), "profile-1", "campaign-1"), result.reconciled_campaigns)

        verified = attribute_ad_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", ("MSKU-A",), ("ASIN-A",))],
            product_costs=[self._line("campaign-1", "19.21", seller_sku="MSKU-A", asin="ASIN-A")],
            campaign_costs=[self._campaign("campaign-1", "19.21")],
            currency_and_date_basis_verified=True,
            mapping_temporal_scope_verified=True,
        )
        self.assertTrue(verified.rows[0].formal_ready)

    def test_conflicting_or_nonunique_product_ids_are_not_attributed(self) -> None:
        result = attribute_ad_costs(
            store_sid=10,
            mappings=[
                AttributionMapping("LOCAL-A", ("MSKU-A",), ("ASIN-A",)),
                AttributionMapping("LOCAL-B", ("MSKU-B",), ("ASIN-B",)),
                AttributionMapping("LOCAL-C", ("MSKU-SHARED",), ()),
                AttributionMapping("LOCAL-D", ("MSKU-SHARED",), ()),
            ],
            product_costs=[
                self._line("campaign-1", "1", seller_sku="MSKU-A", asin="ASIN-B", ad_id="conflict"),
                self._line("campaign-1", "2", seller_sku="MSKU-SHARED", ad_id="shared"),
                self._line("campaign-1", "3", seller_sku="UNKNOWN", asin="UNKNOWN-ASIN", ad_id="unmatched"),
            ],
            campaign_costs=[self._campaign("campaign-1", "6")],
        )
        self.assertEqual([row.local_sku for row in result.rows], ["", "", ""])
        self.assertIn("不同本地 SKU", result.rows[0].status)
        self.assertIn("多个本地 SKU", result.rows[1].status)
        self.assertIn("没有匹配", result.rows[2].status)

    def test_campaign_reconciliation_is_per_campaign_and_does_not_guess_names(self) -> None:
        result = attribute_ad_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", ("MSKU-A",), ("ASIN-A",))],
            product_costs=[
                self._line("campaign-verified", "4.00", seller_sku="MSKU-A", asin="ASIN-A", ad_id="good"),
                self._line("campaign-named-LOCAL-A", "3.00", seller_sku="MSKU-A", asin="ASIN-A", ad_id="bad"),
            ],
            campaign_costs=[
                self._campaign("campaign-verified", "4.00"),
                self._campaign("campaign-named-LOCAL-A", "5.00"),
            ],
        )
        self.assertEqual(result.rows[0].local_sku, "LOCAL-A")
        self.assertTrue(result.rows[0].campaign_reconciled)
        self.assertEqual(result.rows[1].local_sku, "")
        self.assertFalse(result.rows[1].campaign_reconciled)
        self.assertEqual(result.daily_costs()[("SP", date(2026, 8, 29), "local-a")], Decimal("4.00"))

    def test_conflicting_product_rows_for_one_ad_id_block_campaign(self) -> None:
        result = attribute_ad_costs(
            store_sid=10,
            mappings=[
                AttributionMapping("LOCAL-A", ("MSKU-A",), ("ASIN-A",)),
                AttributionMapping("LOCAL-B", ("MSKU-B",), ("ASIN-B",)),
            ],
            product_costs=[
                self._line("campaign-1", "1", seller_sku="MSKU-A", asin="ASIN-A", ad_id="reused-ad"),
                self._line("campaign-1", "2", seller_sku="MSKU-B", asin="ASIN-B", ad_id="reused-ad"),
            ],
            campaign_costs=[self._campaign("campaign-1", "3")],
        )
        self.assertTrue(all(not row.local_sku for row in result.rows))
        self.assertFalse(result.reconciled_campaigns)
        self.assertIn("重复", result.rows[1].status)

    def test_dated_mapping_and_unverified_allocation_fail_closed(self) -> None:
        first = date(2026, 8, 29)
        result = attribute_ad_costs(
            store_sid=10,
            mappings=[
                AttributionMapping("LOCAL-A", ("OLD-MSKU",), (), valid_through=date(2026, 8, 28), validity_verified=True),
                AttributionMapping("LOCAL-A", ("NEW-MSKU",), ("ASIN-A",), valid_from=first, validity_verified=True),
            ],
            product_costs=[
                self._line("campaign-1", "19.21", seller_sku="NEW-MSKU", asin="ASIN-A"),
            ],
            campaign_costs=[self._campaign("campaign-1", "19.21")],
        )
        self.assertEqual(result.rows[0].local_sku, "LOCAL-A")

        unverified = AdCostEvidence(
            "SB", 10, first, "profile-1", "sb-campaign", Decimal("19.21"), "", "", "", "ASIN-A",
            source="verified_asin_allocation", allocation_verified=False,
        )
        blocked = attribute_ad_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", (), ("ASIN-A",))],
            product_costs=[unverified],
            campaign_costs=[CampaignCostEvidence("SB", 10, first, "profile-1", "sb-campaign", Decimal("19.21"))],
        )
        self.assertFalse(blocked.rows[0].local_sku)
        self.assertIn("来源无效", blocked.rows[0].status)

        verified = attribute_ad_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", (), ("ASIN-A",))],
            product_costs=[AdCostEvidence(
                "SB", 10, first, "profile-1", "sb-campaign", Decimal("19.21"),
                asin="ASIN-A", source="verified_asin_allocation", allocation_verified=True,
            )],
            campaign_costs=[CampaignCostEvidence("SB", 10, first, "profile-1", "sb-campaign", Decimal("19.21"))],
        )
        self.assertEqual(verified.rows[0].local_sku, "LOCAL-A")

    def test_exclusive_ad_group_targets_assign_actual_group_spend_only_with_complete_catalog(self) -> None:
        report_day = date(2026, 8, 29)
        campaigns = [CampaignCostEvidence("SB", 10, report_day, "profile-1", "campaign-1", Decimal("19.21"))]
        groups = [
            AdGroupCostEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", Decimal("19.21")),
        ]
        targets = [
            PromotionTargetEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", "ad-a", "", "ASIN-A", identity_valid_for_report=True),
            PromotionTargetEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", "ad-b", "", "ASIN-A2", identity_valid_for_report=True),
        ]
        result = attribute_exclusive_ad_group_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", (), ("ASIN-A", "ASIN-A2"))],
            promoted_targets=targets,
            ad_group_costs=groups,
            campaign_costs=campaigns,
            promotion_catalog_complete=True,
        )
        self.assertEqual(result.rows[0].local_sku, "LOCAL-A")
        self.assertEqual(result.daily_costs()[("SB", report_day, "local-a")], Decimal("19.21"))

        incomplete = attribute_exclusive_ad_group_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", (), ("ASIN-A", "ASIN-A2"))],
            promoted_targets=targets,
            ad_group_costs=groups,
            campaign_costs=campaigns,
            promotion_catalog_complete=False,
        )
        self.assertFalse(incomplete.rows[0].local_sku)
        self.assertIn("完整性未核验", incomplete.rows[0].status)

    def test_shared_ad_group_targets_remain_unattributed(self) -> None:
        report_day = date(2026, 8, 29)
        result = attribute_exclusive_ad_group_costs(
            store_sid=10,
            mappings=[
                AttributionMapping("LOCAL-A", (), ("ASIN-A",)),
                AttributionMapping("LOCAL-B", (), ("ASIN-B",)),
            ],
            promoted_targets=[
                PromotionTargetEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", asin="ASIN-A", identity_valid_for_report=True),
                PromotionTargetEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", asin="ASIN-B", identity_valid_for_report=True),
            ],
            ad_group_costs=[AdGroupCostEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", Decimal("19.21"))],
            campaign_costs=[CampaignCostEvidence("SB", 10, report_day, "profile-1", "campaign-1", Decimal("19.21"))],
            promotion_catalog_complete=True,
        )
        self.assertFalse(result.rows[0].local_sku)
        self.assertIn("推广多个本地 SKU", result.rows[0].status)

    def test_duplicate_promotion_catalog_rows_block_exclusive_group_assignment(self) -> None:
        report_day = date(2026, 8, 29)
        target = PromotionTargetEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", "ad-a", "", "ASIN-A", identity_valid_for_report=True)
        result = attribute_exclusive_ad_group_costs(
            store_sid=10,
            mappings=[AttributionMapping("LOCAL-A", (), ("ASIN-A",))],
            promoted_targets=[target, target],
            ad_group_costs=[AdGroupCostEvidence("SB", 10, report_day, "profile-1", "campaign-1", "group-a", Decimal("19.21"))],
            campaign_costs=[CampaignCostEvidence("SB", 10, report_day, "profile-1", "campaign-1", Decimal("19.21"))],
            promotion_catalog_complete=True,
        )
        self.assertFalse(result.rows[0].local_sku)
        self.assertIn("重复记录", result.rows[0].status)


class McpClientTests(unittest.TestCase):
    def test_streamable_http_init_tools_and_read_only_call(self) -> None:
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request)
            if request.method == "DELETE":
                return httpx.Response(200)
            payload = json.loads(request.content) if request.content else {}
            method = payload.get("method")
            request_id = payload.get("id")
            if method == "initialize":
                return httpx.Response(
                    200,
                    headers={"Content-Type": "application/json", "Mcp-Session-Id": "test-session"},
                    json={"jsonrpc": "2.0", "id": request_id, "result": {"protocolVersion": "2025-03-26", "serverInfo": {"name": "fixture"}}},
                )
            if method == "notifications/initialized":
                return httpx.Response(202)
            if method == "tools/list":
                cursor = payload.get("params", {}).get("cursor")
                tools = ([{"name": "read_report", "description": "read", "annotations": {"readOnlyHint": True}, "inputSchema": {"type": "object"}}]
                         if not cursor else [{"name": "write_report", "annotations": {"readOnlyHint": False}}])
                result = {"tools": tools}
                if not cursor:
                    result["nextCursor"] = "page-2"
                return httpx.Response(200, headers={"Content-Type": "application/json"}, json={"jsonrpc": "2.0", "id": request_id, "result": result})
            if method == "tools/call":
                return httpx.Response(200, headers={"Content-Type": "application/json"}, json={"jsonrpc": "2.0", "id": request_id, "result": {"content": [{"type": "text", "text": "query ok"}]}})
            return httpx.Response(405)

        http = httpx.Client(transport=httpx.MockTransport(handler))
        client = LingxingMcpClient("https://mcp.example.test/mcp", "test-key-never-log", http_client=http)
        self.assertEqual(client.connect()["protocolVersion"], "2025-03-26")
        tools = client.list_tools()
        self.assertEqual([tool.name for tool in tools], ["read_report", "write_report"])
        self.assertTrue(tools[0].read_only)
        self.assertEqual(tools[0].input_schema, {"type": "object"})
        self.assertEqual(tools[0].annotations, {"readOnlyHint": True})
        self.assertEqual(client.call_read_only_tool(tools[0], {})["content"][0]["text"], "query ok")
        with self.assertRaises(McpClientError):
            client.call_read_only_tool(tools[1], {})
        self.assertEqual(requests[0].headers["X-Mcp-Key"], "test-key-never-log")
        self.assertEqual(requests[2].headers["Mcp-Session-Id"], "test-session")
        self.assertEqual(requests[2].headers["MCP-Protocol-Version"], "2025-03-26")
        client.close()
        http.close()

    def test_http_error_does_not_echo_server_body_or_secret(self) -> None:
        def handler(request: httpx.Request) -> httpx.Response:
            if request.method == "DELETE":
                return httpx.Response(200)
            return httpx.Response(403, text="test-key-never-log")

        http = httpx.Client(transport=httpx.MockTransport(handler))
        client = LingxingMcpClient("https://mcp.example.test/mcp", "test-key-never-log", http_client=http)
        with self.assertRaises(McpClientError) as raised:
            client.connect()
        self.assertNotIn("test-key-never-log", str(raised.exception))
        client.close()
        http.close()

    def test_secret_fields_are_redacted_for_display(self) -> None:
        self.assertEqual(redact_sensitive({"X-Mcp-Key": "private", "data": "ok"}), {"X-Mcp-Key": "[已隐藏]", "data": "ok"})
        self.assertEqual(redact_sensitive("echoed key", secret_values=("key",)), "echoed [已隐藏]")

    def test_tool_directory_search_returns_business_schema_even_without_read_only_hint(self) -> None:
        business = {
            "toolId": "get_sku_daily_sales",
            "name": "SKU每日销量与销售额",
            "description": "查询 SKU 每日销量和销售额",
            "inputSchema": {
                "type": "object",
                "properties": {"sku": {"type": "string"}, "start_date": {"type": "string"}},
                "required": ["sku"],
                "additionalProperties": False,
            },
        }
        requests: list[dict] = []

        def handler(request: httpx.Request) -> httpx.Response:
            if request.method == "DELETE":
                return httpx.Response(200)
            payload = json.loads(request.content) if request.content else {}
            method = payload.get("method")
            request_id = payload.get("id")
            if method == "initialize":
                return httpx.Response(200, headers={"Mcp-Session-Id": "s"}, json={"jsonrpc": "2.0", "id": request_id, "result": {"protocolVersion": "2025-03-26"}})
            if method == "notifications/initialized":
                return httpx.Response(202)
            if method == "tools/list":
                return httpx.Response(200, json={"jsonrpc": "2.0", "id": request_id, "result": {"tools": [
                    {"name": "search", "description": "搜索领星业务工具目录并返回对应 Schema", "annotations": {"readOnlyHint": False}, "inputSchema": {"type": "object", "properties": {"keyword": {"type": "string"}}, "additionalProperties": False}},
                    {"name": "action", "description": "执行指定业务工具；调用前必须通过 search 获取 Schema", "annotations": {}, "inputSchema": {"type": "object"}},
                    {"name": "status", "description": "状态", "annotations": {}, "inputSchema": {"type": "object"}},
                ]}})
            if method == "tools/call":
                requests.append(payload)
                return httpx.Response(200, json={"jsonrpc": "2.0", "id": request_id, "result": {"structuredContent": {"tools": [business]}}})
            return httpx.Response(405)

        http = httpx.Client(transport=httpx.MockTransport(handler))
        client = LingxingMcpClient("https://mcp.example.test/mcp", "fixture-key", http_client=http)
        client.connect()
        tools = client.list_tools()
        self.assertEqual(len(tools), 3)
        self.assertTrue(is_directory_search_tool(tools[0]))
        self.assertFalse(tools[1].read_only)
        result = client.call_directory_search(tools[0], {"keyword": "销量"})
        specs = extract_business_tools(result)
        self.assertEqual(len(specs), 1)
        self.assertEqual(specs[0].tool_id, "get_sku_daily_sales")
        self.assertEqual(specs[0].input_schema, business["inputSchema"])
        self.assertEqual(requests[0]["params"]["name"], "search")
        self.assertEqual(requests[0]["params"]["arguments"], {"keyword": "销量"})
        client.close()
        http.close()

    def test_search_tool_without_catalog_contract_is_never_called(self) -> None:
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            if request.method == "DELETE":
                return httpx.Response(200)
            payload = json.loads(request.content) if request.content else {}
            request_id = payload.get("id")
            if payload.get("method") == "initialize":
                return httpx.Response(200, headers={"Mcp-Session-Id": "s"}, json={"jsonrpc": "2.0", "id": request_id, "result": {"protocolVersion": "2025-03-26"}})
            if payload.get("method") == "notifications/initialized":
                return httpx.Response(202)
            if payload.get("method") == "tools/call":
                requests.append(payload)
            return httpx.Response(200, json={"jsonrpc": "2.0", "id": request_id, "result": {"tools": []}})

        http = httpx.Client(transport=httpx.MockTransport(handler))
        client = LingxingMcpClient("https://mcp.example.test/mcp", "fixture-key", http_client=http)
        client.connect()
        unsafe_search = McpTool("search", "查询销售数据", {"type": "object", "properties": {}}, {})
        with self.assertRaises(McpClientError):
            client.call_directory_search(unsafe_search, {})
        self.assertEqual(requests, [])
        client.close()
        http.close()

    def test_action_requires_search_schema_user_grant_and_query_only_candidate(self) -> None:
        business = BusinessTool(
            "query_product_performance_asin_lists",
            "商品表现 ASIN 列表查询",
            "查询商品表现 ASIN 列表",
            {
                "type": "object",
                "properties": {
                    "offset": {"type": "integer"},
                    "length": {"type": "integer"},
                    "search_value": {"type": "array"},
                },
                "required": ["offset", "length", "search_value"],
                "additionalProperties": False,
            },
        )
        search_tool = McpTool(
            "search", "搜索业务工具目录和 Schema", {"type": "object", "properties": {}, "additionalProperties": False}, {}
        )
        action_tool = McpTool(
            "action",
            "执行指定业务工具；调用前必须先通过 search 获取 Schema",
            {
                "type": "object",
                "properties": {
                    "toolId": {"type": "string"},
                    "params": {
                        "type": "object",
                        "properties": {},
                        "additionalProperties": False,
                        "maxProperties": 0,
                    },
                },
                "required": ["toolId", "params"],
                "additionalProperties": False,
            },
            {},
        )
        calls: list[dict] = []

        def handler(request: httpx.Request) -> httpx.Response:
            if request.method == "DELETE":
                return httpx.Response(200)
            payload = json.loads(request.content) if request.content else {}
            method = payload.get("method")
            request_id = payload.get("id")
            if method == "initialize":
                return httpx.Response(200, headers={"Mcp-Session-Id": "s"}, json={"jsonrpc": "2.0", "id": request_id, "result": {"protocolVersion": "2025-03-26"}})
            if method == "notifications/initialized":
                return httpx.Response(202)
            if method == "tools/call":
                params = payload["params"]
                calls.append(params)
                if params["name"] == "search":
                    return httpx.Response(200, json={"jsonrpc": "2.0", "id": request_id, "result": {"structuredContent": {"tools": [{"id": business.tool_id, "name": business.name, "description": business.description, "inputSchema": business.input_schema}]}}})
                return httpx.Response(200, json={"jsonrpc": "2.0", "id": request_id, "result": {"content": [{"type": "text", "text": "query complete"}]}})
            return httpx.Response(200, json={"jsonrpc": "2.0", "id": request_id, "result": {"tools": []}})

        http = httpx.Client(transport=httpx.MockTransport(handler))
        client = LingxingMcpClient("https://mcp.example.test/mcp", "fixture-key", http_client=http)
        client.connect()
        client.call_directory_search(search_tool, {})
        grant = VerifiedReadOnlyGrant(business.tool_id, business.permission_fingerprint)
        valid_arguments = {"offset": 0, "length": 5, "search_value": ["local-test-a"]}
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, business, valid_arguments, VerifiedReadOnlyGrant("other_id", business.permission_fingerprint))
        self.assertEqual(len(calls), 1, "unverified action must not be sent")
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, business, {"length": 5, "search_value": ["local-test-a"]}, grant)
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, business, {"offset": 0, "search_value": ["local-test-a"]}, grant)
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, business, {**valid_arguments, "unexpected": "value"}, grant)
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, business, {**valid_arguments, "offset": "0"}, grant)
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, business, {**valid_arguments, "search_value": [{"sku": "local-test-a"}]}, grant)
        changed_schema = BusinessTool(
            business.tool_id,
            business.name,
            business.description,
            {"type": "object", "properties": {"offset": {"type": "string"}, "length": {"type": "integer"}}, "required": ["offset", "length"], "additionalProperties": False},
        )
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(action_tool, changed_schema, {"offset": 0, "length": 5}, grant)
        bad_outer_schema = McpTool(
            action_tool.name,
            action_tool.description,
            {
                "type": "object",
                "properties": {
                    "toolId": {"type": "integer"},
                    "params": {"type": "object", "properties": {}, "additionalProperties": False},
                },
                "required": ["toolId", "params"],
                "additionalProperties": False,
            },
            {},
        )
        with self.assertRaises(McpClientError):
            client.call_verified_business_query(bad_outer_schema, business, valid_arguments, grant)
        self.assertEqual(len(calls), 1, "invalid params and a changed ID/Schema must not send action")
        result = client.call_verified_business_query(action_tool, business, valid_arguments, grant)
        self.assertIn("content", result)
        self.assertEqual(calls[-1], {"name": "action", "arguments": {"toolId": business.tool_id, "params": valid_arguments}})
        unsafe = BusinessTool("update_listing", "修改Listing", "更新商品信息", business.input_schema)
        self.assertFalse(is_read_only_query_candidate(unsafe))
        camel_case_query = BusinessTool("getSkuDailySales", "SKU销量", "查询 SKU 每日销量", business.input_schema)
        self.assertTrue(is_read_only_query_candidate(camel_case_query))
        operation_schema = BusinessTool(
            "get_report",
            "查询报表",
            "查询报表数据",
            {"type": "object", "properties": {"update_listing": {"type": "boolean"}}},
        )
        self.assertFalse(is_read_only_query_candidate(operation_schema))
        client.close()
        http.close()

    def test_array_without_items_uses_scalar_allowlist_and_keeps_array_constraints(self) -> None:
        schema = {
            "type": "object",
            "properties": {
                "search_value": {
                    "type": "array",
                    "minItems": 1,
                    "maxItems": 3,
                    "uniqueItems": True,
                }
            },
            "required": ["search_value"],
            "additionalProperties": False,
        }
        validate_json_schema(schema, {"search_value": ["local-test-a"]})
        validate_json_schema(schema, {"search_value": ["local-test-a", 0, None]})
        with self.assertRaises(McpClientError):
            validate_json_schema(schema, {"search_value": []})
        with self.assertRaises(McpClientError):
            validate_json_schema(schema, {"search_value": ["local-test-a", "local-test-a"]})
        with self.assertRaises(McpClientError):
            validate_json_schema(schema, {"search_value": ["local-test-a", "local-test-b", "local-test-c", "local-test-d"]})
        with self.assertRaises(McpClientError):
            validate_json_schema(schema, {"search_value": [{"sku": "local-test-a"}]})
        with self.assertRaises(McpClientError):
            validate_json_schema(schema, {"search_value": [["local-test-a"]]})


class LingxingOpenApiTests(unittest.TestCase):
    @staticmethod
    def _seller_record() -> dict[str, object]:
        return {
            "sid": 987654,
            "mid": 3001,
            "name": "Synthetic Store",
            "seller_id": "synthetic-seller-id",
            "account_name": "synthetic-account",
            "seller_account_id": 40001,
            "region": "R_TEST",
            "country": "ZZ",
            "marketplace_id": "MARKETPLACE_TEST",
            "status": 17,
            "has_ads_setting": 23,
        }

    @classmethod
    def _seller_response(cls, data: object | None = None) -> dict[str, object]:
        return {
            "code": 0,
            "message": "success",
            "error_details": [],
            "request_id": "synthetic-request-id",
            "response_time": "2026-01-02T03:04:05Z",
            "data": [cls._seller_record()] if data is None else data,
        }

    def test_seller_response_requires_business_success_code_and_maps_fields(self) -> None:
        payload = self._seller_response()
        sellers = parse_seller_list_response(payload)
        self.assertEqual(len(sellers), 1)
        seller = sellers[0]
        self.assertEqual(seller.sid, 987654)
        self.assertEqual(seller.mid, 3001)
        self.assertEqual(seller.name, "Synthetic Store")
        self.assertEqual(seller.seller_id, "synthetic-seller-id")
        self.assertEqual(seller.account_name, "synthetic-account")
        self.assertEqual(seller.seller_account_id, 40001)
        self.assertEqual(seller.site_label, "ZZ / R_TEST / MARKETPLACE_TEST")
        # These values remain raw integers; no business meaning is inferred.
        self.assertEqual(seller.has_ads_setting, 23)
        self.assertEqual(seller.status, 17)
        self.assertEqual(seller.display_label, "Synthetic Store — ZZ / R_TEST / MARKETPLACE_TEST — SID 987654")
        self.assertNotIn("synthetic-seller-id", repr(seller))
        self.assertNotIn("987654", repr(seller))

    def test_seller_response_accepts_empty_direct_data_array(self) -> None:
        self.assertEqual(parse_seller_list_response(self._seller_response([])), [])

    def test_seller_response_rejects_business_errors_and_wrong_envelope_types(self) -> None:
        with self.assertRaisesRegex(LingxingApiError, "业务 code 9"):
            parse_seller_list_response({**self._seller_response(), "code": 9})
        with self.assertRaisesRegex(LingxingApiError, "integer 类型的业务 code"):
            parse_seller_list_response({**self._seller_response(), "code": "0"})
        with self.assertRaisesRegex(LingxingApiError, "integer 类型的业务 code"):
            parse_seller_list_response({**self._seller_response(), "code": True})
        with self.assertRaisesRegex(LingxingApiError, "data 数组"):
            parse_seller_list_response({key: value for key, value in self._seller_response().items() if key != "data"})
        with self.assertRaisesRegex(LingxingApiError, "data 必须是数组"):
            parse_seller_list_response(self._seller_response({"list": [self._seller_record()]}))
        for field, wrong_value in (
            ("message", 1),
            ("error_details", "none"),
            ("request_id", 12),
            ("response_time", None),
        ):
            with self.subTest(field=field), self.assertRaisesRegex(LingxingApiError, f"{field} 的类型无效"):
                parse_seller_list_response({**self._seller_response(), field: wrong_value})

    def test_seller_response_rejects_missing_or_wrong_record_field_types(self) -> None:
        for field in ("sid", "mid", "seller_account_id", "status", "has_ads_setting"):
            record = self._seller_record()
            record[field] = "not-an-integer"
            with self.subTest(field=field), self.assertRaisesRegex(LingxingApiError, f"字段 {field} 的类型无效"):
                parse_seller_list_response(self._seller_response([record]))
        for field in ("name", "seller_id", "account_name", "region", "country", "marketplace_id"):
            record = self._seller_record()
            record[field] = 123
            with self.subTest(field=field), self.assertRaisesRegex(LingxingApiError, f"字段 {field} 的类型无效"):
                parse_seller_list_response(self._seller_response([record]))
        record = self._seller_record()
        del record["seller_id"]
        with self.assertRaisesRegex(LingxingApiError, "字段 seller_id 的类型无效"):
            parse_seller_list_response(self._seller_response([record]))
        record = self._seller_record()
        record["sid"] = True
        with self.assertRaisesRegex(LingxingApiError, "字段 sid 的类型无效"):
            parse_seller_list_response(self._seller_response([record]))

    @staticmethod
    def _product_record(local_sku: str = "local-test-a", *, offset: int = 0) -> dict[str, object]:
        return {
            "sku": local_sku,
            "price_list": [
                {"local_sku": local_sku, "seller_sku": f"market-test-{offset + 1}", "sid": "987654"}
            ],
            "asins": [{"asin": f"ASINTEST{offset + 1:04d}", "sid": "987654"}],
        }

    def test_long_sales_history_is_split_into_nonoverlapping_92_day_windows(self) -> None:
        start = date(2026, 1, 1)
        end = start + timedelta(days=93)
        windows = split_date_windows(start, end)
        self.assertEqual([(left, right) for left, right in windows], [
            (start, start + timedelta(days=91)),
            (start + timedelta(days=92), end),
        ])
        self.assertTrue(all((right - left).days + 1 <= 92 for left, right in windows))

        mapping_calls: list[tuple[date, date]] = []
        daily_calls: list[tuple[date, date]] = []
        with LingxingOpenApiClient("synthetic-app-id", "synthetic-secret") as client:
            def query_mappings(**kwargs):
                mapping_calls.append((kwargs["start_date"], kwargs["end_date"]))
                return [SkuProductMapping("local-test-a", ("MSKU-A",), ())]

            def query_sales(**kwargs):
                daily_calls.append((kwargs["start_date"], kwargs["end_date"]))
                return []

            with patch.object(client, "query_product_mappings", side_effect=query_mappings), \
                    patch.object(client, "query_daily_sales", side_effect=query_sales):
                result = client.query_daily_sales_history(
                    access_token="synthetic-token",
                    store_sid=987654,
                    local_skus=["local-test-a"],
                    start_date=start,
                    end_date=end,
                )
        self.assertEqual(mapping_calls, windows)
        self.assertEqual(daily_calls, windows)
        self.assertEqual([(window.start_date, window.end_date) for window in result], windows)

    @staticmethod
    def _product_response(records: list[dict[str, object]], *, total: int | None = None) -> dict[str, object]:
        return {
            "code": 0,
            "msg": "success",
            "trace_id": "synthetic-trace-id",
            "data": {"total": len(records) if total is None else total, "list": records},
        }

    @classmethod
    def _product_metrics_record(
        cls,
        local_sku: str = "local-test-a",
        *,
        sid: int | str = 987654,
    ) -> dict[str, object]:
        record = cls._product_record(local_sku)
        record.update({
            "sids": [sid],
            "currency_code": "USD",
            "sessions_total": 5,
            "ad_order_quantity": 2,
            "ad_sales_amount": "18.75",
            "impressions": 250,
            "clicks": 10,
            "ctr": "0.04",
            "ad_cvr": "0.2",
            "avg_star": "4.3",
            "reviews_count": 7,
            "small_cate_rank": [{"category": "Synthetic Category", "rank": 42, "prev_rank": 45}],
            "cate_rank": 120,
            "prev_cate_rank": 130,
            "rank_category": "Synthetic Category",
            "ranking_update_time": None,
            "spend": "12.34",
            "ads_sp_cost": "2.00",
            "shared_ads_sb_cost": "3.47",
            "shared_ads_sbv_cost": "1.00",
            "ads_sd_cost": "5.87",
            "shared_ads_al_cost": "0.00",
            "shared_ads_cc_cost": "0.00",
            "shared_ads_sspaot_cost": "0.00",
            "shared_ads_sar_cost": "0.00",
            "shared_cost_of_advertising": None,
            "adv_rate": "0.25",
        })
        return record

    def test_product_performance_parser_preserves_metrics_and_missing_zero_distinction(self) -> None:
        day = date(2026, 8, 29)
        record = self._product_metrics_record()
        # Synthetic fixture values only; no live account data is used here.
        record["shared_ads_sb_cost"] = "0.00"
        record["shared_ads_sbv_cost"] = None
        del record["ads_sd_cost"]
        total, rows = parse_product_performance_metrics_page(
            self._product_response([record]),
            store_sid=987654,
            requested_skus={"local-test-a"},
            start_date=day,
            end_date=day,
        )
        self.assertEqual(total, 1)
        row = rows[0]
        self.assertIsInstance(row, ProductPerformanceMetrics)
        self.assertEqual(row.seller_skus, ("market-test-1",))
        self.assertEqual(row.asins, ("ASINTEST0001",))
        self.assertEqual((row.start_date, row.end_date), (day, day))
        self.assertEqual(row.currency_code, "USD")
        self.assertEqual(row.values["shared_ads_sb_cost"], Decimal("0.00"))
        self.assertIn("shared_ads_sb_cost", row.present_fields)
        self.assertIsNone(row.values["shared_ads_sbv_cost"])
        self.assertIn("shared_ads_sbv_cost", row.present_fields)
        self.assertIsNone(row.values["ads_sd_cost"])
        self.assertNotIn("ads_sd_cost", row.present_fields)
        self.assertEqual(row.values["ad_sales_amount"], Decimal("18.75"))
        self.assertEqual(row.values["small_cate_rank"][0], {
            "category": "Synthetic Category", "rank": 42, "prev_rank": 45,
        })

    def test_product_performance_parser_accepts_documented_null_rank_and_rejects_scope_errors(self) -> None:
        day = date(2026, 8, 29)
        record = self._product_metrics_record()
        record["small_cate_rank"] = None
        _, rows = parse_product_performance_metrics_page(
            self._product_response([record]),
            store_sid=987654,
            requested_skus={"local-test-a"},
            start_date=day,
            end_date=day,
        )
        self.assertIsNone(rows[0].values["small_cate_rank"])

        for changes, message in (
            ({"sids": [987654, 123456]}, "店铺范围"),
            ({"sku": "other-sku"}, "未请求的 local SKU"),
            ({"spend": "not-a-number"}, "数值格式无效"),
            ({"small_cate_rank": {"rank": 3}}, "small_cate_rank 应为数组或 null"),
        ):
            bad_record = {**self._product_metrics_record(), **changes}
            with self.subTest(message=message), self.assertRaisesRegex(LingxingApiError, message):
                parse_product_performance_metrics_page(
                    self._product_response([bad_record]),
                    store_sid=987654,
                    requested_skus={"local-test-a"},
                    start_date=day,
                    end_date=day,
                )

        with self.assertRaisesRegex(LingxingApiError, "业务 code 7"):
            parse_product_performance_metrics_page(
                {**self._product_response([]), "code": 7},
                store_sid=987654,
                requested_skus={"local-test-a"},
                start_date=day,
                end_date=day,
            )

    def test_product_performance_parser_resolves_msku_summary_to_requested_local_sku(self) -> None:
        day = date(2026, 8, 29)
        first = self._product_metrics_record()
        first["sku"] = None
        second = self._product_metrics_record()
        second["sku"] = None
        second["price_list"] = [
            {"local_sku": "local-test-a", "seller_sku": "market-test-2", "sid": "987654"}
        ]
        total, rows = parse_product_performance_metrics_page(
            self._product_response([first, second]),
            store_sid=987654,
            requested_skus={"local-test-a"},
            start_date=day,
            end_date=day,
            summary_field="msku",
        )
        self.assertEqual(total, 2)
        self.assertEqual([(row.local_sku, row.summary_field, row.summary_key) for row in rows], [
            ("local-test-a", "msku", "market-test-1"),
            ("local-test-a", "msku", "market-test-2"),
        ])

        ambiguous = self._product_metrics_record()
        ambiguous["sku"] = None
        ambiguous["price_list"] = [
            {"local_sku": "local-test-a", "seller_sku": "market-test-1", "sid": "987654"},
            {"local_sku": "local-test-b", "seller_sku": "market-test-2", "sid": "987654"},
        ]
        with self.assertRaisesRegex(LingxingApiError, "无法唯一匹配一个本店"):
            parse_product_performance_metrics_page(
                self._product_response([ambiguous]),
                store_sid=987654,
                requested_skus={"local-test-a", "local-test-b"},
                start_date=day,
                end_date=day,
                summary_field="msku",
            )

    def test_product_performance_single_day_query_uses_local_sku_sku_summary_and_paginates(self) -> None:
        app_id = "synthetic-app-id"
        token = "synthetic-token"
        day = date(2026, 8, 29)
        sku_names = ["local-test-a", "local-test-b"]
        requests: list[tuple[httpx.Request, dict[str, object]]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            requests.append((request, body))
            offset = body["offset"]
            record = self._product_metrics_record(sku_names[offset]) if offset < len(sku_names) else None
            records = [record] if record else []
            return httpx.Response(200, json=self._product_response(records, total=len(sku_names)))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api.PRODUCT_PERFORMANCE_PAGE_LENGTH", 1), \
                patch("amazon_weekly_report.providers.lingxing_api._PRODUCT_PERFORMANCE_LAST_RESPONSE_AT", None), \
                patch("amazon_weekly_report.providers.lingxing_api.time.time", return_value=1639734344), \
                patch("amazon_weekly_report.providers.lingxing_api.time.monotonic", return_value=100.0), \
                patch("amazon_weekly_report.providers.lingxing_api.time.sleep"):
            with LingxingOpenApiClient(app_id, "synthetic-secret", http_client=http) as client:
                rows = client.query_product_performance_single_day(
                    access_token=token,
                    store_sid=987654,
                    local_skus=sku_names,
                    event_date=day,
                )
        http.close()

        self.assertEqual([row.local_sku for row in rows], sku_names)
        self.assertEqual([body["offset"] for _, body in requests], [0, 1])
        self.assertEqual(len(requests), 2)
        for request, body in requests:
            self.assertEqual(request.method, "POST")
            self.assertEqual(request.url.path, "/bd/productPerformance/openApi/asinList")
            self.assertEqual(body["search_field"], "local_sku")
            self.assertEqual(body["search_value"], sku_names)
            self.assertEqual(body["summary_field"], "sku")
            self.assertEqual(body["sid"], "987654")
            self.assertEqual(body["start_date"], day.isoformat())
            self.assertEqual(body["end_date"], day.isoformat())
            self.assertEqual(body["purchase_status"], 0)
            self.assertIs(body["is_recently_enum"], False)
            self.assertNotIn("currency_code", body)
            query = parse_qs(request.url.query.decode("ascii"))
            self.assertEqual(query["access_token"], [token])
            self.assertEqual(query["timestamp"], ["1639734344"])
            self.assertEqual(query["sign"], [signed_post_query_parameters({
                "access_token": token,
                "app_key": app_id,
                "timestamp": 1639734344,
            }, body, app_id)["sign"]])

    def test_product_performance_cost_diagnostic_sends_three_dates_and_middle_day_msku(self) -> None:
        token = "synthetic-token"
        start = date(2026, 8, 28)
        end = date(2026, 8, 30)
        requests: list[dict[str, object]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            requests.append(body)
            record = self._product_metrics_record()
            if body["summary_field"] == "msku":
                record["sku"] = None
            return httpx.Response(200, json=self._product_response([record]))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api._PRODUCT_PERFORMANCE_LAST_RESPONSE_AT", None), \
                patch("amazon_weekly_report.providers.lingxing_api.time.time", return_value=1639734344), \
                patch("amazon_weekly_report.providers.lingxing_api.time.monotonic", return_value=100.0), \
                patch("amazon_weekly_report.providers.lingxing_api.time.sleep"):
            with LingxingOpenApiClient("synthetic-app-id", "synthetic-secret", http_client=http) as client:
                rows = client.query_product_performance_cost_diagnostic(
                    access_token=token,
                    store_sid=987654,
                    local_sku="local-test-a",
                    start_date=start,
                    end_date=end,
                )
        http.close()

        self.assertEqual(len(requests), 4)
        self.assertEqual(
            [(body["start_date"], body["summary_field"]) for body in requests],
            [("2026-08-28", "sku"), ("2026-08-29", "sku"),
             ("2026-08-30", "sku"), ("2026-08-29", "msku")],
        )
        self.assertTrue(all(body["start_date"] == body["end_date"] for body in requests))
        self.assertTrue(all(body["search_field"] == "local_sku" for body in requests))
        self.assertTrue(all(body["search_value"] == ["local-test-a"] for body in requests))
        self.assertEqual([row.summary_field for row in rows], ["sku", "sku", "sku", "msku"])
        self.assertEqual([row.start_date for row in rows], [
            start, start + timedelta(days=1), end, start + timedelta(days=1),
        ])

        with LingxingOpenApiClient("synthetic-app-id", "synthetic-secret") as client:
            with self.assertRaisesRegex(LingxingApiError, "连续三天"):
                client.query_product_performance_cost_diagnostic(
                    access_token=token,
                    store_sid=987654,
                    local_sku="local-test-a",
                    start_date=start,
                    end_date=start + timedelta(days=1),
                )

    @staticmethod
    def _daily_record(
        day: str,
        metric_value: int | float,
        *,
        seller_sku: str = "market-test-1",
        asin: str = "ASINTEST0001",
    ) -> dict[str, object]:
        return {
            "sid": 987654,
            "r_date": day,
            "currency_code": "USD",
            "seller_sku": seller_sku,
            "asin": asin,
            "product_name": "Synthetic Product",
            "map_value": metric_value,
        }

    @staticmethod
    def _daily_response(records: list[dict[str, object]]) -> dict[str, object]:
        return {
            "code": 0,
            "message": "success",
            "error_details": [],
            "request_id": "synthetic-request-id",
            "response_time": "2026-01-02T03:04:05Z",
            "data": records,
        }

    def test_product_performance_parser_resolves_local_sku_to_store_msku_and_asin_candidates(self) -> None:
        payload = self._product_response([self._product_record()])
        total, mappings = parse_product_performance_page(
            payload,
            store_sid=987654,
            requested_skus={"local-test-a"},
        )
        self.assertEqual(total, 1)
        self.assertEqual(
            mappings,
            [SkuProductMapping("local-test-a", ("market-test-1",), ("ASINTEST0001",))],
        )

    def test_product_performance_parser_rejects_bad_envelope_and_untyped_mapping(self) -> None:
        good = self._product_response([self._product_record()])
        with self.assertRaisesRegex(LingxingApiError, "业务 code 7"):
            parse_product_performance_page(
                {**good, "code": 7}, store_sid=987654, requested_skus={"local-test-a"}
            )
        bad_row = self._product_record()
        bad_row["price_list"] = "not-an-array"
        with self.assertRaisesRegex(LingxingApiError, "price_list/asins 必须是数组"):
            parse_product_performance_page(
                self._product_response([bad_row]),
                store_sid=987654,
                requested_skus={"local-test-a"},
            )
        with self.assertRaisesRegex(LingxingApiError, "price_list.sid 的类型无效"):
            parse_product_performance_page(
                self._product_response([{
                    **self._product_record(),
                    "price_list": [{"local_sku": "local-test-a", "seller_sku": "market-test-1", "sid": True}],
                }]),
                store_sid=987654,
                requested_skus={"local-test-a"},
            )

    def test_daily_parser_preserves_explicit_zero_and_rejects_bad_site_date_or_types(self) -> None:
        day = date(2026, 1, 2)
        parsed = parse_asin_daily_page(
            self._daily_response([self._daily_record(day.isoformat(), 0)]),
            store_sid=987654,
            event_date=day,
            metric_type=2,
        )
        self.assertEqual(parsed[0].metric_value, Decimal("0"))
        self.assertEqual(parsed[0].event_date, day)
        with self.assertRaisesRegex(LingxingApiError, "站点日期不一致"):
            parse_asin_daily_page(
                self._daily_response([self._daily_record("2026-01-03", 1)]),
                store_sid=987654,
                event_date=day,
                metric_type=2,
            )
        with self.assertRaisesRegex(LingxingApiError, "不是整数"):
            parse_asin_daily_page(
                self._daily_response([self._daily_record(day.isoformat(), 1.5)]),
                store_sid=987654,
                event_date=day,
                metric_type=3,
            )
        with self.assertRaisesRegex(LingxingApiError, "SID 不匹配"):
            parse_asin_daily_page(
                self._daily_response([{**self._daily_record(day.isoformat(), 1), "sid": 1}]),
                store_sid=987654,
                event_date=day,
                metric_type=2,
            )

    def test_product_and_daily_queries_send_documented_parameters_and_keep_missing_distinct_from_zero(self) -> None:
        app_id = "ak_0123456789abc"
        token = "fixture-token"
        first_day = date(2026, 1, 2)
        second_day = date(2026, 1, 3)
        requests: list[tuple[httpx.Request, dict[str, object]]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            requests.append((request, body))
            if request.url.path == "/bd/productPerformance/openApi/asinList":
                return httpx.Response(200, json=self._product_response([self._product_record()]))
            self.assertEqual(request.url.path, "/erp/sc/data/sales_report/asinDailyLists")
            event_date = body["event_date"]
            metric_type = body["type"]
            if event_date == first_day.isoformat():
                values = {1: 12.34, 2: 2, 3: 1}
                records = [self._daily_record(event_date, values[metric_type])]
            elif metric_type == 2:
                # A returned zero is present data; omitted type 1/type 3 records
                # stay None in the normalized result.
                records = [self._daily_record(event_date, 0)]
            else:
                records = []
            return httpx.Response(200, json=self._daily_response(records))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api._PRODUCT_PERFORMANCE_LAST_RESPONSE_AT", None), \
                patch("amazon_weekly_report.providers.lingxing_api.time.time", return_value=1639734344), \
                patch("amazon_weekly_report.providers.lingxing_api.time.monotonic", return_value=100.0), \
                patch("amazon_weekly_report.providers.lingxing_api.time.sleep"):
            with LingxingOpenApiClient(app_id, "fixture-app-secret", http_client=http) as client:
                mappings = client.query_product_mappings(
                    access_token=token,
                    store_sid=987654,
                    local_skus=["local-test-a"],
                    start_date=first_day,
                    end_date=second_day,
                )
                rows = client.query_daily_sales(
                    access_token=token,
                    store_sid=987654,
                    mappings=mappings,
                    start_date=first_day,
                    end_date=second_day,
                )
        http.close()

        self.assertEqual(len(requests), 7)
        product_request, product_body = requests[0]
        self.assertEqual(product_request.method, "POST")
        self.assertEqual(product_request.url.path, "/bd/productPerformance/openApi/asinList")
        self.assertEqual(product_body["search_field"], "local_sku")
        self.assertEqual(product_body["search_value"], ["local-test-a"])
        self.assertEqual(product_body["length"], 10_000)
        self.assertEqual(product_body["sid"], "987654")
        self.assertEqual(product_body["summary_field"], "sku")
        self.assertIs(product_body["is_recently_enum"], False)
        self.assertEqual(product_body["start_date"], first_day.isoformat())
        self.assertEqual(product_body["end_date"], second_day.isoformat())
        self.assertEqual([request.url.path for request, _ in requests[1:]], [
            "/erp/sc/data/sales_report/asinDailyLists"
        ] * 6)
        for request, body in requests[1:]:
            self.assertEqual(body["sid"], 987654)
            self.assertEqual(body["asin_type"], 2)
            self.assertIn(body["event_date"], (first_day.isoformat(), second_day.isoformat()))
            self.assertIn(body["type"], (1, 2, 3))
            self.assertEqual(body["length"], 1000)
            query = parse_qs(request.url.query.decode("ascii"))
            self.assertEqual(query["access_token"], [token])
            self.assertEqual(query["timestamp"], ["1639734344"])
            self.assertEqual(query["sign"], [signed_post_query_parameters({
                "access_token": token,
                "app_key": app_id,
                "timestamp": 1639734344,
            }, body, app_id)["sign"]])

        self.assertEqual(len(rows), 2)
        first, second = rows
        self.assertEqual((first.units_sold, first.order_count, first.sales_amount), (2, 1, Decimal("12.34")))
        self.assertEqual(first.event_date, first_day)
        self.assertEqual(first.asin, "ASINTEST0001")
        self.assertEqual((second.units_sold, second.order_count, second.sales_amount), (0, None, None))
        self.assertEqual(second.event_date, second_day)
        self.assertEqual(second.asin, "ASINTEST0001")

    def test_product_mapping_paginates_using_offset_and_rejects_uncovered_pages(self) -> None:
        requested = ["local-test-a", "local-test-b"]
        offsets: list[int] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            offsets.append(body["offset"])
            record_index = body["offset"]
            records = [
                self._product_record(requested[record_index], offset=record_index)
            ] if record_index < len(requested) else []
            return httpx.Response(200, json=self._product_response(records, total=len(requested)))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api.PRODUCT_PERFORMANCE_PAGE_LENGTH", 1), \
                patch("amazon_weekly_report.providers.lingxing_api._PRODUCT_PERFORMANCE_LAST_RESPONSE_AT", None), \
                patch("amazon_weekly_report.providers.lingxing_api.time.monotonic", return_value=100.0), \
                patch("amazon_weekly_report.providers.lingxing_api.time.sleep"):
            with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
                mappings = client.query_product_mappings(
                    access_token="fixture-token",
                    store_sid=987654,
                    local_skus=requested,
                    start_date=date(2026, 1, 1),
                    end_date=date(2026, 1, 2),
                )
        http.close()
        self.assertEqual(offsets, [0, 1])
        self.assertEqual([mapping.local_sku for mapping in mappings], requested)

    def test_product_mapping_rejects_duplicate_local_sku_across_pages(self) -> None:
        record = self._product_record()

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            return httpx.Response(200, json=self._product_response([record], total=2))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api.PRODUCT_PERFORMANCE_PAGE_LENGTH", 1), \
                patch("amazon_weekly_report.providers.lingxing_api._PRODUCT_PERFORMANCE_LAST_RESPONSE_AT", None), \
                patch("amazon_weekly_report.providers.lingxing_api.time.monotonic", return_value=100.0), \
                patch("amazon_weekly_report.providers.lingxing_api.time.sleep"):
            with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
                with self.assertRaisesRegex(LingxingApiError, "重复 local SKU"):
                    client.query_product_mappings(
                        access_token="fixture-token",
                        store_sid=987654,
                        local_skus=["local-test-a", "local-test-b"],
                        start_date=date(2026, 1, 1),
                        end_date=date(2026, 1, 2),
                    )
        http.close()

    def test_daily_statistics_paginates_serially_and_stops_on_rate_limit(self) -> None:
        day = date(2026, 1, 2)
        offsets: list[tuple[int, int]] = []

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            metric_type = body["type"]
            offset = body["offset"]
            offsets.append((metric_type, offset))
            page_rows = [
                self._daily_record(day.isoformat(), 3, seller_sku="market-other"),
                self._daily_record(day.isoformat(), {1: 12.34, 2: 2, 3: 1}[metric_type]),
            ]
            records = page_rows[offset:offset + 1]
            return httpx.Response(200, json=self._daily_response(records))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api.ASIN_DAILY_PAGE_LENGTH", 1):
            with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
                rows = client.query_daily_sales(
                    access_token="fixture-token",
                    store_sid=987654,
                    mappings=[SkuProductMapping("local-test-a", ("market-test-1",), ("ASINTEST0001",))],
                    start_date=day,
                    end_date=day,
                )
        http.close()
        self.assertEqual(offsets, [(metric, offset) for metric in (1, 2, 3) for offset in (0, 1, 2)])
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0].units_sold, rows[0].order_count, rows[0].sales_amount), (2, 1, Decimal("12.34")))

        limited_http = httpx.Client(
            transport=httpx.MockTransport(lambda _request: httpx.Response(429, headers={"Retry-After": "9"}))
        )
        with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=limited_http) as client:
            with self.assertRaisesRegex(LingxingApiError, "等待 9 秒"):
                client.query_daily_sales(
                    access_token="fixture-token",
                    store_sid=987654,
                    mappings=[SkuProductMapping("local-test-a", ("market-test-1",), ())],
                    start_date=day,
                    end_date=day,
                )
        limited_http.close()

    def test_daily_statistics_rejects_duplicate_msku_across_pages(self) -> None:
        day = date(2026, 1, 2)

        def handler(request: httpx.Request) -> httpx.Response:
            body = json.loads(request.content.decode("utf-8"))
            return httpx.Response(200, json=self._daily_response([self._daily_record(day.isoformat(), 1)]))

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with patch("amazon_weekly_report.providers.lingxing_api.ASIN_DAILY_PAGE_LENGTH", 1):
            with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
                with self.assertRaisesRegex(LingxingApiError, "重复 MSKU"):
                    client.query_daily_sales(
                        access_token="fixture-token",
                        store_sid=987654,
                        mappings=[SkuProductMapping("local-test-a", ("market-test-1",), ())],
                        start_date=day,
                        end_date=day,
                    )
        http.close()

    def test_signature_matches_reference_algorithm_vector(self) -> None:
        app_id = "ak_0123456789abc"
        parameters = {
            "timestamp": 1639734344,
            "access_token": "fixture-token",
            "app_key": app_id,
        }
        self.assertEqual(
            canonical_signing_string(parameters),
            "access_token=fixture-token&app_key=ak_0123456789abc&timestamp=1639734344",
        )
        self.assertEqual(
            generate_signature(parameters, app_id),
            "DGIcRRQN/VQlrOtSjtCmntt5i7TsON7XPBbSVCNFij5IwVhSPzVcd0sPeLkAl3rO",
        )

    def test_signature_serialization_boundaries_fail_closed(self) -> None:
        self.assertEqual(
            canonical_signing_string({"z": "  value  ", "a": "", "enabled": True, "count": 5}),
            "a=&count=5&enabled=true&z=value",
        )
        for value in (None, 1.5, ["x"], {"x": "y"}):
            with self.subTest(value_type=type(value).__name__):
                with self.assertRaises(LingxingSignatureError):
                    canonical_signing_string({"value": value})
        with self.assertRaises(LingxingSignatureError):
            canonical_signing_string({"签名": "value"})
        with self.assertRaises(LingxingSignatureError):
            canonical_signing_string({"sign": "self-reference"})
        with self.assertRaises(LingxingSignatureError):
            generate_signature({"timestamp": 1}, "short")

    def test_signed_post_serializes_arrays_as_compact_json_for_sdk_signature(self) -> None:
        app_id = "ak_0123456789abc"
        auth = {"access_token": "fixture-token", "app_key": app_id, "timestamp": 1639734344}
        body = {
            "search_value": ["local sku", "本地SKU"],
            "sid": "987654",
            "offset": 0,
        }
        signed = signed_post_query_parameters(auth, body, app_id)
        canonical_parameters = {
            **auth,
            "search_value": '["local sku","本地SKU"]',
            "sid": "987654",
            "offset": 0,
        }
        self.assertEqual(signed["sign"], generate_signature(canonical_parameters, app_id))
        self.assertEqual(set(signed), {"access_token", "app_key", "timestamp", "sign"})
        with self.assertRaises(LingxingSignatureError):
            signed_post_query_parameters(auth, {"timestamp": 1}, app_id)

    def test_query_encoding_matches_java_urlencoder_for_utf8_and_base64(self) -> None:
        encoded = encode_query_parameters({"value": "a b+/%~=中", "sign": "A+/="})
        self.assertEqual(encoded, "value=a+b%2B%2F%25%7E%3D%E4%B8%AD&sign=A%2B%2F%3D")

    def test_list_sellers_sends_sdk_signed_get_and_parses_business_response(self) -> None:
        app_id = "ak_0123456789abc"
        token = "fixture-token"
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request)
            return httpx.Response(
                200,
                json={
                    "code": 0,
                    "message": "success",
                    "error_details": [],
                    "request_id": "synthetic-request-id",
                    "response_time": "2026-01-02T03:04:05Z",
                    "data": [
                        self._seller_record(),
                    ],
                },
            )

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient(app_id, "fixture-app-secret", http_client=http) as client:
            with patch("amazon_weekly_report.providers.lingxing_api.time.time", return_value=1639734344):
                response = client.list_sellers(token)

        request = requests[0]
        self.assertEqual(request.method, "GET")
        self.assertEqual(request.url.path, "/erp/sc/data/seller/lists")
        self.assertEqual(request.content, b"")
        self.assertNotIn("fixture-app-secret", str(request.url))
        wire_query = request.url.query.decode("ascii")
        self.assertIn("%2F", wire_query, "Base64 slash must be URL encoded")
        query = parse_qs(wire_query, keep_blank_values=True)
        self.assertEqual(query["access_token"], [token])
        self.assertEqual(query["app_key"], [app_id])
        self.assertEqual(query["timestamp"], ["1639734344"])
        self.assertEqual(
            query["sign"],
            [generate_signature({"access_token": token, "app_key": app_id, "timestamp": 1639734344}, app_id)],
        )
        self.assertEqual(len(response), 1)
        self.assertEqual(response[0].name, "Synthetic Store")
        self.assertEqual(response[0].sid, 987654)
        self.assertEqual(response[0].has_ads_setting, 23)
        http.close()

    def test_list_sellers_rejects_http_200_business_error(self) -> None:
        payload = self._seller_response()
        payload["code"] = 41
        http = httpx.Client(
            transport=httpx.MockTransport(lambda _request: httpx.Response(200, json=payload))
        )
        with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
            with self.assertRaisesRegex(LingxingApiError, "业务 code 41"):
                client.list_sellers("fixture-token")
        http.close()

    def test_list_sellers_http_error_does_not_echo_token_or_signature(self) -> None:
        token = "fixture-access-token"
        http = httpx.Client(
            transport=httpx.MockTransport(
                lambda _request: httpx.Response(403, text=f"denied {token} fake-signature")
            )
        )
        with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
            with self.assertRaises(LingxingApiError) as raised:
                client.list_sellers(token)
        self.assertNotIn(token, str(raised.exception))
        self.assertNotIn("fake-signature", str(raised.exception))
        http.close()

    def test_list_sellers_rejects_non_success_redirect_response(self) -> None:
        http = httpx.Client(
            transport=httpx.MockTransport(lambda _request: httpx.Response(302, json={"code": 0, "data": []}))
        )
        with LingxingOpenApiClient("ak_0123456789abc", "fixture-app-secret", http_client=http) as client:
            with self.assertRaisesRegex(LingxingApiError, "HTTP 302"):
                client.list_sellers("fixture-token")
        http.close()

    def test_token_exchange_uses_multipart_and_hides_token(self) -> None:
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request)
            return httpx.Response(
                200,
                json={"code": 0, "data": {"access_token": "fixture-token-do-not-display", "expires_in": 3600}},
            )

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("fixture-app-id", "fixture-app-secret", http_client=http) as client:
            result = client.request_access_token()
        request = requests[0]
        self.assertEqual(request.url.path, "/api/auth-server/oauth/access-token")
        self.assertIn("multipart/form-data; boundary=", request.headers["Content-Type"])
        body = request.content.decode("utf-8")
        self.assertIn('name="appId"', body)
        self.assertIn("fixture-app-id", body)
        self.assertIn('name="appSecret"', body)
        self.assertIn("fixture-app-secret", body)
        self.assertEqual(result.expires_in_seconds, 3600)
        self.assertNotIn("fixture-token-do-not-display", repr(result))
        self.assertNotIn("fixture-token-do-not-display", result.safe_status())
        http.close()

    def test_access_token_cache_reuses_only_tokens_with_expiry(self) -> None:
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request)
            token_number = len(requests)
            return httpx.Response(
                200,
                json={"code": 0, "data": {"access_token": f"fixture-token-{token_number}", "expires_in": 3600}},
            )

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("fixture-app-id", "fixture-app-secret", http_client=http) as client:
            first = client.get_access_token()
            cached = client.get_access_token()
            self.assertIs(first, cached)
            self.assertEqual(len(requests), 1)

            refreshed = client.get_access_token(force_refresh=True)
            self.assertEqual(refreshed.token, "fixture-token-2")
            self.assertEqual(len(requests), 2)

            client.clear_token_cache()
            after_clear = client.get_access_token()
            self.assertEqual(after_clear.token, "fixture-token-3")
            self.assertEqual(len(requests), 3)
        http.close()

    def test_access_token_without_expiry_is_not_cached(self) -> None:
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request)
            return httpx.Response(200, json={"code": 0, "data": {"access_token": f"fixture-token-{len(requests)}"}})

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("fixture-app-id", "fixture-app-secret", http_client=http) as client:
            first = client.get_access_token()
            second = client.get_access_token()
        self.assertEqual((first.token, second.token), ("fixture-token-1", "fixture-token-2"))
        self.assertEqual(len(requests), 2)
        http.close()

    def test_access_token_near_expiry_is_refreshed(self) -> None:
        requests: list[httpx.Request] = []

        def handler(request: httpx.Request) -> httpx.Response:
            requests.append(request)
            expires_in = 1 if len(requests) == 1 else 3600
            return httpx.Response(
                200,
                json={"code": 0, "data": {"access_token": f"fixture-token-{len(requests)}", "expires_in": expires_in}},
            )

        http = httpx.Client(transport=httpx.MockTransport(handler))
        with LingxingOpenApiClient("fixture-app-id", "fixture-app-secret", http_client=http) as client:
            first = client.get_access_token()
            second = client.get_access_token()
        self.assertEqual((first.token, second.token), ("fixture-token-1", "fixture-token-2"))
        self.assertEqual(len(requests), 2)
        http.close()

    def test_unverified_sku_sales_query_fails_closed_without_sending_request(self) -> None:
        requests: list[httpx.Request] = []
        http = httpx.Client(transport=httpx.MockTransport(lambda request: requests.append(request) or httpx.Response(200)))
        with LingxingOpenApiClient("fixture-app-id", "fixture-app-secret", http_client=http) as client:
            with self.assertRaises(LingxingApiNotReady):
                client.query_sku_sales(
                    access_token="fixture-token",
                    seller_id="fixture-seller",
                    sku="local-test-a",
                    start_date="2026-08-01",
                    end_date="2026-08-10",
                )
        self.assertEqual(requests, [])
        http.close()

    def test_token_http_error_does_not_echo_credentials(self) -> None:
        http = httpx.Client(transport=httpx.MockTransport(lambda _request: httpx.Response(403, text="fixture-app-secret")))
        with LingxingOpenApiClient("fixture-app-id", "fixture-app-secret", http_client=http) as client:
            with self.assertRaises(LingxingApiError) as raised:
                client.request_access_token()
        self.assertNotIn("fixture-app-secret", str(raised.exception))
        http.close()


class OpenApiPreviewTests(unittest.TestCase):
    class _Value:
        def __init__(self) -> None:
            self.value = ""

        def set(self, value: str) -> None:
            self.value = value

    class _Tree:
        def __init__(self) -> None:
            self.rows: list[tuple[object, ...]] = []

        def insert(self, _parent: str, _index: str, *, values: tuple[object, ...]) -> None:
            self.rows.append(values)

    def test_gui_preview_formats_missing_and_explicit_zero_without_claiming_historical_first_sale(self) -> None:
        app = object.__new__(WeeklyReportApp)
        app.report_mapping_status = self._Value()
        app.report_query_status = self._Value()
        app.report_sales_tree = self._Tree()
        day = date(2026, 1, 2)
        result = {
            "_mappings": [SkuProductMapping("local-test-a", ("market-test-1",), ("ASINTEST0001",))],
            "_daily_rows": [
                SkuDailySales(
                    sid=987654,
                    local_sku="local-test-a",
                    seller_sku="market-test-1",
                    asin="ASINTEST0001",
                    event_date=day,
                    currency_code="USD",
                    units_sold=0,
                    order_count=None,
                    sales_amount=None,
                )
            ],
            "_requested_skus": ["local-test-a"],
            "_unresolved_skus": [],
            "_start_date": day,
            "_end_date": day,
            "_store_name": "Synthetic Store",
        }
        app._apply_daily_sales_preview(result)
        self.assertEqual(app.report_sales_tree.rows[0][4:8], ("0", "", "", "USD"))
        self.assertIn("ASIN 候选未假定与某个 MSKU 一一对应", app.report_mapping_status.value)
        self.assertIn("所选范围未发现明确正销量记录", app.report_query_status.value)
        self.assertIn("不是已核验的历史首单日期", app.report_query_status.value)

    def test_product_performance_preview_distinguishes_zero_null_and_absent_fields(self) -> None:
        row = ProductPerformanceMetrics(
            store_sid=987654,
            local_sku="local-test-a",
            seller_skus=("market-test-1",),
            asins=("ASINTEST0001",),
            start_date=date(2026, 8, 29),
            end_date=date(2026, 8, 29),
            currency_code="USD",
            values={"shared_ads_sb_cost": Decimal("0"), "shared_ads_sbv_cost": None},
            present_fields=frozenset({"shared_ads_sb_cost", "shared_ads_sbv_cost"}),
        )
        self.assertEqual(_product_performance_value(row, "shared_ads_sb_cost"), "0")
        self.assertIn("null", _product_performance_value(row, "shared_ads_sbv_cost"))
        self.assertEqual(_product_performance_value(row, "ads_sd_cost"), "字段未返回")

    def test_cost_diagnostic_preview_shows_date_dimension_and_caveat_without_report_write(self) -> None:
        app = object.__new__(WeeklyReportApp)
        app.product_performance_status = self._Value()
        app.product_performance_tree = self._Tree()
        start = date(2026, 8, 28)
        rows = [
            ProductPerformanceMetrics(
                store_sid=987654,
                local_sku="local-test-a",
                seller_skus=(summary_key,),
                asins=("ASINTEST0001",),
                start_date=event_date,
                end_date=event_date,
                currency_code="USD",
                values={"shared_ads_sb_cost": Decimal("3.20"), "currency_code": "USD"},
                present_fields=frozenset({"shared_ads_sb_cost", "currency_code"}),
                summary_field=summary_field,
                summary_key=summary_key,
            )
            for event_date, summary_field, summary_key in (
                (start, "sku", "local-test-a"),
                (start + timedelta(days=1), "sku", "local-test-a"),
                (start + timedelta(days=2), "sku", "local-test-a"),
                (start + timedelta(days=1), "msku", "market-test-1"),
            )
        ]
        app._apply_product_performance_cost_diagnostic({
            "_product_performance_rows": rows,
            "_product_performance_skus": ["local-test-a"],
            "_product_performance_start_date": start,
            "_product_performance_end_date": start + timedelta(days=2),
            "_product_performance_store": "Synthetic Store",
            "_product_performance_sid": 987654,
        })

        self.assertEqual(len(app.product_performance_tree.rows), 4 * 7)
        self.assertEqual(app.product_performance_tree.rows[0][:4], (
            "2026-08-28", "sku", "local-test-a", "local-test-a",
        ))
        self.assertEqual(app.product_performance_tree.rows[-7][:2], ("2026-08-29", "msku"))
        self.assertIn("不能单独证明", app.product_performance_status.value)
        self.assertNotIn("19.21", str(app.product_performance_tree.rows))


if __name__ == "__main__":
    unittest.main()
